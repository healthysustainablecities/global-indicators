"""Export indicator-dashboard layers from a region database.

Produces the per-layer newline-delimited GeoJSON inputs for building PMTiles
vector-tile archives (via tippecanoe, run separately by ``build_tiles.sh``),
together with the three JSON files the dashboard reads:

    manifest.json     scales, layers, bounding box, region-level values
    stats.json        per (scale, column) weighted summaries, and the share of
                      population in each colour class
    indicators.json   the faceted indicator vocabulary and its prose, plus the
                      colour classes and targets every column is drawn with

This generalises ``_export_validation_tiles.py`` (the cycling validation site's
exporter) in three ways: every aggregation scale is exported rather than only
the population grid; the indicator vocabulary is derived from the region's own
accessibility configuration and data dictionary rather than a hardcoded cycling
cross product; and the explanatory text the viewer needs is emitted as data
instead of being hardcoded in its JavaScript.

Three types of dashboard can be exported (``dashboard.type``, or ``kind``):

    general     every configured indicator, grouped by theme, without
                composite indices
    composite   one composite index (``dashboard.index``) as a dashboard of
                its own: its scores, its variants, and its indicators' raw
                values, with ``distributions.json`` comparing the regions'
                results as smoothed histograms
    combined    both: every indicator, and the index as a theme of its own,
                so the viewer's theme selector moves between them

The regular 100 m grid is also written as a raster of cell values
(``<layer>.cells.i32`` and one ``<layer>__<group>.f32`` per tile group, listed
in the scale's ``raster`` entry), from which the viewer draws a smoothed
surface on request.

Each scale is written as one layer per *tile group* of columns
(``scale_<key>__<group>.geojsonl``; see :func:`tile_groups`), so that no
feature carries so many values that tippecanoe must drop features from the
zoomed out tiles.

Usage (inside the ghsci container), either from the region object:

    r.export_dashboard()
    r.export_dashboard(outdir=..., scales=['grid'], layers=False)
    r.export_dashboard(kind='general')

or on the commandline, from the process folder:

    /env/bin/python subprocesses/_export_dashboard.py <config.yml> [outdir]
        [--scales=a,b,c] [--type=general|composite] [--no-layers]

Default outdir is /tmp/dashboard_export/<slug>/ (copy out with docker cp).
"""

import csv
import json
import math
import os
import re
import shutil
import sys

if __name__ == '__main__':
    # usage examples give configuration paths relative to the process
    # folder; as a module this leaves the caller's directory alone
    os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import ghsci  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from _accessibility_spec import (  # noqa: E402
    AVOID,
    DEFAULT_DESTINATIONS,
    STANDARD_SET,
    accessibility_config,
    activity_centre_definitions,
    combined_access_sets,
    diversity_sets,
    effective_config,
    spec_thresholds,
)
from _cycling_accessibility import (  # noqa: E402
    DMGAP_INFIX,
    MEASURE_ORDER,
    MEASURES,
)
from _cycling_lts_network import _data_path  # noqa: E402
from _utils import slugify  # noqa: E402
from data_dictionary import (  # noqa: E402
    URBAN_HEAT_STRUCTURE,
    describe_units,
    describe_variable,
)

# (parent index, sub-indicator names) for the urban heat composite, so the
# picker lists each sub-index ahead of the sub-indicators feeding it.
URBAN_HEAT_STRUCTURE_PARENTS = URBAN_HEAT_STRUCTURE

# ---------------------------------------------------------------------------
# Networks.  Walking is the implicit base network -- it has no infix because its
# columns carry their own 'walk' word -- and every cycling measure contributes a
# network whose column infix is defined once, in _cycling_accessibility.MEASURES.
# 'dmgap' is not a routable network but the paired contrast between two of them;
# it is offered as a network so the viewer can present it in the same control.
WALK = 'walk'
# A network has two names: the one a dropdown shows ("Caminando"), and the one
# a sentence needs ("using the *pedestrian* network").  The second is the
# `phrase`; without it every composed sentence reads "using the Walking
# network", which is why they used to be written as "red: caminando" instead.
NETWORKS = {
    WALK: {
        'infix': None,
        'mode': 'walk',
        'label': {'es': 'Caminando', 'en': 'Walking'},
        'phrase': {'es': 'peatonal', 'en': 'pedestrian'},
        'description': {
            'es': 'Red peatonal completa, sin restricción de tráfico.',
            'en': 'The full pedestrian network, with no traffic restriction.',
        },
    },
}
# Spanish for the cycling measures.  MEASURES is shared with the analysis and
# the reports, which are English, so the translation lives here rather than
# being pushed back into the pipeline's own vocabulary.
NETWORK_ES = {
    'lts1': {
        'label': 'Ruta de bajo estrés (sólo LTS 1)',
        'phrase': 'ciclista de muy bajo estrés',
        'description': 'geométrica, totalmente LTS 1, permite desmontar',
    },
    'low_stress_ride': {
        'label': 'Pedaleo de bajo estrés, sin desmontar (LTS 1–2)',
        'phrase': 'ciclista de bajo estrés, sin desmontar',
        'description': 'geométrica, totalmente LTS ≤ 2, pedaleada de principio a fin',
    },
    'low_stress': {
        'label': 'Ruta de bajo estrés (LTS 1–2)',
        'phrase': 'ciclista de bajo estrés',
        'description': 'geométrica, totalmente LTS ≤ 2, permite empujar la bicicleta',
    },
    'danger_weighted': {
        'label': 'Ruta con penalización por estrés',
        'phrase': 'ciclista con penalización por estrés',
        'description': 'red ciclable completa, penalizando los tramos de mayor estrés',
    },
}
# English phrases for the same, since MEASURES supplies only the long label
NETWORK_EN_PHRASE = {
    'lts1': 'very low stress cycling',
    'low_stress_ride': 'low-stress riding',
    'low_stress': 'low-stress cycling',
    'danger_weighted': 'stress-weighted cycling',
}
for _key in MEASURE_ORDER:
    _m = MEASURES[_key]
    _es = NETWORK_ES.get(_key, {})
    NETWORKS[_key] = {
        'infix': _m['infix'],
        'mode': 'cycle',
        'label': {'en': _m['label'], **({'es': _es['label']} if _es else {})},
        'phrase': {
            'en': NETWORK_EN_PHRASE.get(_key, _m['label'].lower()),
            **({'es': _es['phrase']} if _es.get('phrase') else {}),
        },
        'description': {
            'en': _m['description'],
            **({'es': _es['description']} if _es else {}),
        },
    }
NETWORKS['dmgap'] = {
    'infix': DMGAP_INFIX,
    'mode': 'cycle',
    'label': {
        'es': 'Dependencia de desmontar',
        'en': 'Dismount dependence',
    },
    'phrase': {
        'es': 'que depende de desmontar',
        'en': 'dismount-dependent',
    },
    'description': {
        'en': (
            'The share of access that exists only because the rider may get '
            'off and walk the bike.'
        ),
    },
}
CYCLE_NETWORKS = [k for k in NETWORKS if k not in (WALK, 'dmgap')]

# Spanish labels for the fixed vocabulary.  Indicator descriptions themselves
# come from the generated data dictionary (English); where no Spanish string is
# supplied the viewer falls back to English, so the gap stays visible.
ES = {
    'Daily essential services': 'Servicios esenciales cotidianos',
    'Health, education and care': 'Salud, educación y cuidados',
    'Community, culture and recreation': 'Comunidad, cultura y recreación',
    'Employment': 'Empleo',
    'Transport': 'Transporte',
    'Open space': 'Espacio abierto',
    'Other': 'Otros',
}


# ---------------------------------------------------------------------------
# Configuration


def dashboard_config(r):
    """The region's ``dashboard`` block, with defaults filled in."""
    config = dict((r.config or {}).get('dashboard') or {})
    config.setdefault('slug', slugify(r.name))
    label = config.get('label') or r.name
    config['label'] = label if isinstance(label, dict) else {'en': label}
    # the application's own title, distinct from the dataset's label
    title = config.get('title') or config['label']
    config['title'] = title if isinstance(title, dict) else {'en': title}
    config.setdefault('regions', {})
    config.setdefault('scale_labels', {})
    return config


def _labels(value, fallback):
    """Normalise a label to an ``{es, en}`` mapping."""
    if isinstance(value, dict):
        return {k: v for k, v in value.items() if v}
    return {'en': value or fallback}


# ---------------------------------------------------------------------------
# Scales
#
# A "scale" is one set of areas the indicators have been aggregated to: the
# population grid, each configured custom aggregation, and the study region as a
# whole.  _12_aggregation writes each to ``indicators_<key>`` and keeps the
# imported source boundaries in ``agg_<key>``; the two differ, because areas that
# received no source units are deleted from the indicator table
# (_12_aggregation.py:570).  Exporting from a LEFT JOIN of the boundaries onto the
# indicators is what keeps those areas visible as "no data" rather than as holes.

GRID_SCALE = 'grid'


def _sql_key(name):
    return name.replace(' ', '_').lower()


def discover_scales(r, config):
    """Ordered scale definitions available in this region's database."""
    tables = set(r.get_tables())
    scales = {}

    grid_table = r.config.get('grid_summary')
    if grid_table in tables:
        resolution = r.config['population']['resolution']
        scales[GRID_SCALE] = {
            'key': GRID_SCALE,
            'table': grid_table,
            'boundaries': grid_table,
            'id': 'grid_id',
            'weight': 'pop_est',
            'source': 'grid',
            # the configured label wins here as it does for every custom
            # aggregation; the derived one is only a fallback
            'label': (
                _labels(
                    config['scale_labels'].get(GRID_SCALE),
                    f'{resolution} grid',
                )
                if config['scale_labels'].get(GRID_SCALE)
                else {
                    'es': f'Cuadrícula de {resolution}',
                    'en': f'{resolution} grid',
                }
            ),
        }

    for agg, spec in (r.config.get('custom_aggregations') or {}).items():
        key = _sql_key(agg)
        table, boundaries = f'indicators_{key}', f'agg_{key}'
        if table not in tables:
            continue
        scales[key] = {
            'key': key,
            'table': table,
            'boundaries': boundaries if boundaries in tables else table,
            'id': _sql_key(spec.get('id') or 'ogc_fid'),
            'weight': spec.get('weight'),
            'source': spec.get('aggregation_source'),
            'keep_columns': [
                _sql_key(c)
                for c in str(spec.get('keep_columns') or '').split(',')
                if c.strip()
            ],
            'label': _labels(config['scale_labels'].get(agg), agg),
        }

    requested = config.get('scales')
    if requested:
        missing = [s for s in requested if s not in scales]
        if missing:
            print(f'  ! configured scales not found, skipping: {missing}')
        return [scales[s] for s in requested if s in scales]
    return list(scales.values())


# ---------------------------------------------------------------------------
# Column canonicalisation
#
# The same indicator does not always carry the same column name at every scale.
# The columns propagated from the sample points (pct_access_walk_*,
# avg_cycle_dist_*, urban_heat_*, local_*) do -- they are added to every scale by
# the same pass -- but the columns produced by the weighted aggregation itself
# take a prefix that depends on how the area was weighted: 'pop_' where it was
# weighted by population estimate, none where it was not.  So an AGEB carries
# pop_pct_access_500m_fresh_food_market_score and a Condesa lot carries
# pct_access_500m_fresh_food_market_score for the same quantity.
#
# The exporter resolves each to one canonical (unprefixed) name and renames on
# the way out, so the viewer's paint expressions are scale-independent.  Which
# physical column supplied it is recorded per scale, because it changes what the
# number means and the info panel has to say so.
WEIGHTED_PREFIXES = ('pop_est_', 'pop_')
# canonical name -> the population-weighted variants that may stand in for it
CANONICAL_ALIASES = {
    'local_walkability': ('pop_walkability',),
    'local_daily_living': ('pop_daily_living',),
    'local_nh_population_density': ('pop_nh_pop_density',),
    'local_nh_intersection_density': ('pop_nh_intersection_density',),
}


def table_columns(r, table):
    """Ordered column names of a table."""
    return list(
        r.get_df(
            f"""SELECT column_name FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = '{table}'
                ORDER BY ordinal_position""",
        )['column_name'],
    )


def canonical_columns(r, table):
    """Map canonical indicator names to the physical column carrying them.

    Returns ``{canonical: (physical, weighting)}`` where weighting is 'weighted'
    when the value came from a population-weighted aggregation column, 'direct'
    when the column already carries the canonical name.
    """
    present = table_columns(r, table)
    resolved = {}
    for column in present:
        resolved.setdefault(column, (column, 'direct'))
    for column in present:
        for prefix in WEIGHTED_PREFIXES:
            if column.startswith(prefix):
                canonical = column[len(prefix) :]
                if canonical not in resolved:
                    resolved[canonical] = (column, 'weighted')
                break
    for canonical, aliases in CANONICAL_ALIASES.items():
        if canonical in resolved:
            continue
        for alias in aliases:
            if alias in present:
                resolved[canonical] = (alias, 'weighted')
                break
    return resolved


# ---------------------------------------------------------------------------
# Column algebra
#
# One place where the naming conventions of _pedestrian_accessibility and
# _cycling_accessibility (as carried through to the aggregated scales by
# _12_aggregation) are expressed.  The viewer never rebuilds a column name: it
# reads the ones emitted here, so a convention can change in one place.


def access_column(network, name, distance):
    if network == WALK:
        return f'pct_access_walk_{name}_{distance}m'
    return f'pct_access_cycle_{NETWORKS[network]["infix"]}{name}_{distance}m'


def beyond_column(name, distance):
    """Disamenity polarity: the share living beyond the threshold."""
    return f'pct_beyond_walk_{name}_{distance}m'


def distance_column(network, name):
    if network == WALK:
        return f'avg_walk_dist_{name}'
    if network == 'dmgap':
        # not a distance to the destination but the extra riding needed to
        # avoid the links a rider would otherwise have to walk
        return f'avg_cycle_extra_{DMGAP_INFIX}{name}'
    return f'avg_cycle_dist_{NETWORKS[network]["infix"]}{name}'


def count_column(set_name, group, distance):
    return f'avg_count_walk_{set_name}__{group}_{distance}m'


def diversity_column(set_name, distance):
    return f'avg_diversity_walk_{set_name}_{distance}m'


def richness_column(set_name, distance):
    return f'avg_richness_walk_{set_name}_{distance}m'


# ---------------------------------------------------------------------------
# Indicator vocabulary

HIGHER, LOWER = 'higher_is_better', 'lower_is_better'

# Measure kinds.  'banded' indicators exist at several distance thresholds at
# once and are drawn with the nested-isochrone colour scheme and its clickable
# distance-band legend; 'continuous' indicators take a single value per area and
# are drawn with a class-banded sequential ramp.
BANDED, CONTINUOUS = 'banded', 'continuous'

MEASURE_META = {
    'access': {
        'kind': BANDED,
        'direction': HIGHER,
        'units': 'percent',
        'label': {'es': 'Acceso (%)', 'en': 'Access (%)'},
    },
    'beyond': {
        'kind': BANDED,
        'direction': HIGHER,
        'units': 'percent',
        'label': {'es': 'Más allá de (%)', 'en': 'Beyond (%)'},
    },
    'distance': {
        'kind': CONTINUOUS,
        'direction': LOWER,
        'units': 'metres',
        'label': {
            'es': 'Distancia promedio a lo más cercano (m)',
            'en': 'Average distance to closest (m)',
        },
    },
    'dmextra': {
        'kind': CONTINUOUS,
        'direction': LOWER,
        'units': 'metres',
        'label': {
            'es': 'Desvío adicional para evitar desmontar (m)',
            'en': 'Extra riding to avoid dismounting (m)',
        },
    },
    'count': {
        'kind': CONTINUOUS,
        'direction': HIGHER,
        'units': 'count',
        'label': {
            'es': 'Número promedio alcanzable',
            'en': 'Average number reachable',
        },
    },
    'diversity': {
        'kind': CONTINUOUS,
        'direction': HIGHER,
        'units': 'index 0-1',
        'label': {'es': 'Diversidad', 'en': 'Diversity'},
    },
    'richness': {
        'kind': CONTINUOUS,
        'direction': HIGHER,
        'units': 'index 0-1',
        'label': {'es': 'Riqueza', 'en': 'Richness'},
    },
    'value': {
        'kind': CONTINUOUS,
        'direction': HIGHER,
        'units': '',
        'label': {'es': 'Valor', 'en': 'Value'},
    },
}

# ---------------------------------------------------------------------------
# The sentences the viewer composes.
#
# These are emitted as data rather than written into the viewer's JavaScript, so
# that the wording can be revised in the site's own ``data/<slug>/text.json``
# -- edit, refresh, done -- without a code change, a re-export or a rebuild.
# Anything set there overrides what is here.
#
# Placeholders, substituted from the current selection:
#
#   {network}       the network as a sentence names it ("pedestrian")
#   {networkLabel}  the network as the dropdown names it ("Walking")
#   {distance}      the selected distance, e.g. "500 m"
#   {bands}         every distance offered, "500 m / 1000 m / 1500 m"
#   {indicator}     the indicator's own name
#   {description}   its data-dictionary description
#
# ``showing`` is the plain sentence under the indicator name in the sidebar;
# ``notes`` is the "how to read this" paragraph in the info panel, which is
# about how to read a *kind* of number rather than any one variable.
DEFAULT_TEXT = {
    'showing': {
        'access': {
            'es': (
                'Porcentaje de la población con acceso dentro de {distance}, '
                'usando la red {network}.'
            ),
            'en': (
                'Percentage of the population with access within {distance}, '
                'using the {network} network.'
            ),
        },
        'beyond': {
            'es': (
                'Porcentaje de la población que vive más allá de {distance} '
                'del destino, usando la red {network}. Aquí estar lejos es lo '
                'deseable.'
            ),
            'en': (
                'Percentage of the population living beyond {distance} of the '
                'destination, using the {network} network. Here, being far '
                'away is the good outcome.'
            ),
        },
        'distance': {
            'es': (
                'Distancia promedio a lo más cercano, usando la red {network}.'
            ),
            'en': (
                'Average distance to the closest, using the {network} network.'
            ),
        },
        'dmextra': {
            'es': (
                'Metros adicionales de pedaleo necesarios para evitar los '
                'tramos que de otro modo habría que caminar empujando la '
                'bicicleta.'
            ),
            'en': (
                'Extra metres of riding needed to avoid the links a rider '
                'would otherwise have to walk the bike along.'
            ),
        },
        'count': {
            'es': (
                'Número promedio de establecimientos de este subtipo '
                'alcanzables dentro de {distance}.'
            ),
            'en': (
                'Average number of establishments of this sub-type reachable '
                'within {distance}.'
            ),
        },
        'diversity': {
            'es': (
                'Qué tan uniformemente se reparte entre subtipos lo que es '
                'alcanzable dentro de {distance}: 1 es un reparto '
                'perfectamente uniforme.'
            ),
            'en': (
                'How evenly what is reachable within {distance} is spread '
                'across sub-types: 1 is a perfectly even spread.'
            ),
        },
        'richness': {
            'es': (
                'Proporción de los subtipos configurados que son alcanzables '
                'dentro de {distance} — cuántas clases distintas, no cuántos '
                'establecimientos.'
            ),
            'en': (
                'The share of configured sub-types reachable within '
                '{distance} — how many distinct kinds, not how many '
                'establishments.'
            ),
        },
        'value': {
            'es': '{description}',
            'en': '{description}',
        },
    },
    'notes': {
        'access': {
            'es': (
                'El valor es el porcentaje de los puntos de muestreo de un '
                'área que alcanzan el destino dentro de la distancia '
                'seleccionada, ponderado por población. Cada distancia se '
                'mide por separado, y el acceso a una distancia corta implica '
                'acceso a una más larga: por eso las cifras de la tabla '
                'crecen de izquierda a derecha.'
            ),
            'en': (
                'The value is the percentage of an area’s sample points that '
                'reach the destination within the selected distance, weighted '
                'by population. Each distance is measured separately, and '
                'access within a shorter distance implies access within a '
                'longer one — which is why the figures in the table '
                'rise from left to right.'
            ),
        },
        'beyond': {
            'es': (
                'Este destino se mide con polaridad inversa: es una '
                'desamenidad, y lo que se reporta es la proporción que vive '
                'más allá de la distancia. Un valor más alto es mejor.'
            ),
            'en': (
                'This destination is measured with the opposite polarity: it '
                'is a disamenity, and what is reported is the share living '
                'beyond the distance. A higher value is better.'
            ),
        },
        'distance': {
            'es': (
                'Distancia promedio por la red hasta el destino más cercano. '
                'Se censura en la distancia configurada más amplia, por lo '
                'que un área sin nada dentro de ella no tiene valor en lugar '
                'de una distancia inventada.'
            ),
            'en': (
                'Average network distance to the nearest destination. It is '
                'censored at the largest configured distance, so an area with '
                'nothing within that distance has no value rather than an '
                'invented one.'
            ),
        },
        'dmextra': {
            'es': (
                'Metros adicionales de pedaleo necesarios para evitar los '
                'tramos que de otro modo habría que caminar empujando la '
                'bicicleta. Se promedia sólo donde ese rodeo tiene un costo.'
            ),
            'en': (
                'Extra metres of riding needed to avoid the links a rider '
                'would otherwise have to walk the bike along. Averaged only '
                'where that detour costs anything.'
            ),
        },
        'count': {
            'es': (
                'Número promedio de establecimientos de este subtipo '
                'alcanzables dentro de la distancia. Un conteo no es una '
                'proporción de la población: no está escalado a porcentaje.'
            ),
            'en': (
                'Average number of establishments of this sub-type reachable '
                'within the distance. A count is not a share of a population, '
                'so it is not scaled to a percentage.'
            ),
        },
        'diversity': {
            'es': (
                'Entropía de Shannon normalizada de los subtipos alcanzables: '
                'qué tan uniformemente se reparte lo alcanzable entre los '
                'subtipos configurados. Cinco panaderías y cinco '
                'establecimientos distintos puntúan igual en acceso, y no '
                'deberían.'
            ),
            'en': (
                'Normalised Shannon entropy of the reachable sub-types: how '
                'evenly what is reachable is spread across the configured '
                'sub-types. Five bakeries and five different kinds of shop '
                'score the same for access, and should not.'
            ),
        },
        'richness': {
            'es': (
                'Proporción de los subtipos configurados de un conjunto que '
                'son alcanzables — cuántas clases distintas, no cuántos '
                'establecimientos.'
            ),
            'en': (
                'The share of a set’s configured sub-types that are reachable '
                '— how many distinct kinds, not how many establishments.'
            ),
        },
        'value': {
            'es': (
                'Un único valor por área. Consulte el diccionario de datos '
                'para su definición completa y sus unidades.'
            ),
            'en': (
                'A single value per area. See the data dictionary for its '
                'full definition and units.'
            ),
        },
    },
    'scale_note': {
        'es': (
            'La escala de agregación cambia el resultado: no todas cubren el '
            'mismo terreno. Sólo la cuadrícula y la región cubren toda el '
            'área de estudio.'
        ),
        'en': (
            'The aggregation scale changes the result: they do not all cover '
            'the same ground. Only the grid and the region cover the whole '
            'study area.'
        ),
    },
    # The method of a composite index, as the conceptual model's panel
    # summarises it.  Prose only: the formulas and the full references cited
    # here are drawn by the viewer, being the same in every language.
    # Placeholders: {index} (its name), {domains} (their names), {n} (the
    # number of indicators), {thresholds} (the soft thresholds used), {k} (the
    # soft threshold slope).  'normalise' describes the AMPI and
    # 'normalise_mpi' the classic MPI; the viewer shows the one the index uses.
    'methods': {
        'title': {
            'es': 'Cómo se calcula el índice',
            'en': 'How the index is calculated',
        },
        'intro': {
            'es': (
                '{index} resume {n} indicadores estimados para cada punto de '
                'muestra de la red peatonal, agrupados en los dominios del '
                'modelo conceptual: {domains}. Sigue el enfoque del Índice de '
                'Entornos Vivibles (Urban Liveability Index; Higgs et al. '
                '2019): un índice compuesto de Mazziotta-Pareto, parcialmente '
                'no compensatorio, en el que un buen desempeño en unos '
                'aspectos no compensa del todo un mal desempeño en otros.'
            ),
            'en': (
                '{index} summarises {n} indicators estimated for every sample '
                'point on the pedestrian network, grouped by the domains of '
                'the conceptual model: {domains}. It follows the approach of '
                'the Urban Liveability Index (Higgs et al. 2019): a '
                'Mazziotta-Pareto composite index, which is partially '
                'non-compensatory, so that doing well on some aspects cannot '
                'wholly make up for doing badly on others.'
            ),
        },
        'threshold': {
            'es': (
                'La distancia d al destino más cercano se convierte primero en '
                'una puntuación de acceso suave entre 0 y 1, con el umbral '
                'suave de Higgs et al. (2019): una función logística del '
                'umbral t ({thresholds}) con pendiente k = {k}. Vale 0,5 en el '
                'umbral, casi 1 muy por debajo de él (0,99 en d = 0) y casi 0 '
                'muy por encima (0,007 en d = 2t), y 0 donde no se encontró '
                'nada dentro de la distancia buscada. A diferencia de un corte '
                'binario, una pequeña diferencia de distancia cerca del umbral '
                'produce una pequeña diferencia de puntuación.'
            ),
            'en': (
                'The distance d to the nearest destination is first turned '
                'into a soft access score between 0 and 1, using the soft '
                'threshold of Higgs et al. (2019): a logistic function of the '
                'threshold t ({thresholds}) with slope k = {k}. It is 0.5 at '
                'the threshold, close to 1 well within it (0.99 at d = 0) and '
                'close to 0 well beyond it (0.007 at d = 2t), and 0 where '
                'nothing was found within the distance searched. Unlike a '
                'binary cut-off, a small difference in distance near the '
                'threshold makes only a small difference to the score.'
            ),
        },
        # shown only where an indicator is scored by a ladder of distances;
        # {steps} lists each such indicator's ladder
        'steps': {
            'es': (
                'Algunas distancias se puntúan en cambio con una escala '
                'escalonada publicada: {steps}. Los indicadores que ya son '
                'proporciones, diversidades o índices entran tal cual.'
            ),
            'en': (
                'Some distances are scored instead by a published ladder: '
                '{steps}. Indicators that are already proportions, diversity '
                'scores or indices enter as they are.'
            ),
        },
        'normalise': {
            'es': (
                'Cada indicador se reescala con la normalización del Índice '
                'de Mazziotta-Pareto Ajustado (AMPI; Mazziotta y Pareto '
                '2018), invertida donde un valor menor es mejor, de modo que '
                'un valor mayor siempre indica un entorno más vivible, y se '
                'expresa en puntos por encima o por debajo de la referencia. '
                'La referencia Ref es el promedio del indicador en todos los '
                'puntos de muestra de la región de estudio, y recibe 0. Las '
                'metas Min y Max se centran en Ref y distan entre sí tanto '
                'como el rango observado del indicador, así que ese rango '
                'abarca 60 puntos (de −30 a +30 cuando el promedio queda en '
                'el centro). Un indicador con una escala propia, como un índice '
                'de 0 a 100, puede usarla como metas; entonces sus valores '
                'observados abarcan menos de 60 puntos. Las metas usadas se '
                'registran: puntuar un análisis posterior con las mismas '
                'metas hace comparables las puntuaciones a lo largo del '
                'tiempo.'
            ),
            'en': (
                'Each indicator is re-scaled using the normalisation of the '
                'Adjusted Mazziotta-Pareto Index (AMPI; Mazziotta and Pareto '
                '2018), reversed where lower values are better so that a '
                'higher score always means more liveable, and expressed as '
                'points above or below the reference. The reference Ref is '
                'the indicator\'s mean across all the study region\'s sample '
                'points, and scores 0. The goalposts Min and Max are centred '
                'on Ref and as far apart as the indicator\'s observed range, '
                'so that range spans 60 points (−30 to +30 when the mean lies '
                'midway). An indicator with a scale of its own, '
                'such as an index from 0 to 100, may use that scale as its '
                'goalposts instead, in which case its observed values span '
                'fewer than 60 points. The goalposts used are recorded: '
                'scoring a later analysis against the same goalposts makes '
                'scores comparable over time.'
            ),
        },
        'normalise_mpi': {
            'es': (
                'Cada indicador se estandariza con media 100 y desviación '
                'estándar 10, como en el Índice de Mazziotta-Pareto (MPI; '
                'De Muro, Mazziotta y Pareto 2011) y en el Índice de '
                'Entornos Vivibles (Higgs et al. 2019), invertido donde un '
                'valor menor es mejor. Las puntuaciones son relativas a esta '
                'área de estudio en este momento.'
            ),
            'en': (
                'Each indicator is standardised to a mean of 100 and a '
                'standard deviation of 10, as in the Mazziotta-Pareto Index '
                '(MPI; De Muro, Mazziotta and Pareto 2011) and the Urban '
                'Liveability Index (Higgs et al. 2019), reversed where lower '
                'values are better. Scores are relative to this study area at '
                'this time.'
            ),
        },
        'aggregate': {
            'es': (
                'Dentro de cada dominio se calculan la media ponderada M y la '
                'desviación estándar S de las puntuaciones de sus indicadores '
                '(con pesos w que suman 1: iguales, salvo que se configure '
                'otra cosa o que un indicador se comparta con otros dominios, '
                'en cuyo caso cuenta sólo su parte). La puntuación es M − S·cv, '
                'donde cv = S/M es su '
                'coeficiente de variación: la penalización S·cv crece con el '
                'desequilibrio entre indicadores, de modo que un perfil '
                'equilibrado puntúa más que uno desigual con la misma media '
                '(De Muro, Mazziotta y Pareto 2011; forma ponderada de '
                'Mazziotta y Pareto 2022). El índice se calcula igual a '
                'partir de las puntuaciones de los dominios, sin volver a '
                'normalizarlas, así que cada dominio pesa lo mismo sea cual '
                'sea su número de indicadores. Como cv divide entre la media, '
                'estos cálculos se hacen con la referencia en 100, donde toda '
                'puntuación es positiva, y el resultado se expresa luego '
                'respecto de la referencia (restando 100).'
            ),
            'en': (
                'Within each domain, the weighted mean M and standard '
                'deviation S of its indicators\' scores are found (with '
                'weights w that sum to 1: equal, unless configured otherwise '
                'or an indicator is shared with other domains, when it counts '
                'only its share). '
                'The score is M − S·cv, where cv = S/M is their coefficient '
                'of variation: the penalty S·cv grows with the imbalance '
                'between indicators, so that a balanced profile scores higher '
                'than an uneven one with the same mean (De Muro, Mazziotta '
                'and Pareto 2011; weighted form, Mazziotta and Pareto 2022). '
                'The index is calculated the same way from the domain scores, '
                'without normalising them again, so each domain carries the '
                'same weight whatever its number of indicators. As cv divides '
                'by the mean, these are calculated with the reference at 100, '
                'where every score is positive, and the result is then '
                'expressed relative to the reference (less 100).'
            ),
        },
        # shown only where an indicator counts towards more than one domain
        'shared': {
            'es': (
                'Algunos indicadores atañen a más de un dominio del modelo '
                'conceptual. Un indicador así cuenta en cada dominio al que '
                'pertenece con una parte de su peso (un tercio en cada uno de '
                'tres), de modo que sus partes suman el peso de cualquier otro '
                'indicador; su puntuación es la misma en cada uno. Como cada '
                'dominio pesa lo mismo en el índice sea cual sea su contenido, '
                'el peso efectivo de un indicador en el índice depende aún de '
                'cuántos otros comparten sus dominios: aquí, de {effective_min} '
                'a {effective_max} del nivel medio del índice. Los indicadores '
                'compartidos también hacen que las puntuaciones de sus dominios '
                'varíen juntas, lo que reduce la penalización por '
                'desequilibrio entre dominios (no dentro de ellos).'
            ),
            'en': (
                'Some indicators bear on more than one domain of the '
                'conceptual model. Such an indicator counts a share of its '
                'weight in each domain it belongs to (a third in each of '
                'three), so that its shares add up to the weight of any other '
                'indicator; its score is the same in each. Because every '
                'domain carries the same weight in the index whatever it '
                'holds, an indicator\'s effective weight in the index still '
                'depends on how many others share its domains: here from '
                '{effective_min} to {effective_max} of the index\'s mean '
                'level. Indicators shared between domains also make their '
                'scores move together, which reduces the penalty for '
                'imbalance between domains (though not within them).'
            ),
        },
        'scale': {
            'es': (
                'El índice se calcula en cada punto de muestra y luego se '
                'promedia: a cada celda de la cuadrícula como la media de sus '
                'puntos, y a áreas mayores y a la región como promedios, por '
                'lo general ponderados por población. El índice de un área es, por tanto, el promedio '
                'de los índices de sus puntos, no el índice de sus valores '
                'promedio, y su penalización es el desequilibrio promedio en '
                'sus puntos.'
            ),
            'en': (
                'The index is calculated at each sample point and then '
                'averaged: to each grid cell as the mean of its points, and '
                'to larger areas and the region as averages, usually '
                'weighted by population. An '
                'area\'s index is therefore the average of its points\' '
                'indices, not the index of its average values, and its '
                'penalty is the average imbalance at its points.'
            ),
        },
        'reading': {
            'es': (
                'En este tablero los valores se muestran como puntos por '
                'encima o por debajo de la referencia (0): un lugar donde cada indicador es igual a su promedio '
                'en los puntos de muestra de la región de estudio. Como la '
                'penalización se resta en todas partes y las áreas se ponderan '
                'por población, el índice de la propia región no es '
                'exactamente 0. Cada indicador se observa a través de un '
                'enfoque (proximidad, accesibilidad, densidad...).'
            ),
            'en': (
                'This dashboard shows values as points above or below the '
                'reference (0): a place where every '
                'indicator equals its average over the study region\'s sample '
                'points. Because the penalty is subtracted everywhere, and '
                'areas are weighted by population, the region\'s own index is '
                'not exactly 0. Each indicator is seen through a lens '
                '(proximity, accessibility, density...).'
            ),
        },
        # Walkability and its attenuation by thermal comfort, the one setting
        # a composite dashboard offers.  Shown only for an index with variants.
        'walkability': {
            'es': (
                'La caminabilidad suma los puntajes z del acceso a la vida '
                'diaria (un mercado de alimentos frescos, una tienda de '
                'conveniencia y el transporte público, dentro de 300 m), la '
                'densidad de población y la densidad de intersecciones en el '
                'vecindario caminable de 1 km (Frank et al. 2010), tomados en '
                'todos los puntos de muestra. Por defecto se atenúa por el '
                'confort térmico diurno: su rango percentil se multiplica por '
                '1 − λ·h, donde h es el UTCI diurno medio reescalado entre 0 '
                '(a {lowValue}) y 1 '
                '(a {highValue}), y λ = {attenuation}. El calor '
                'reduce lo caminable que es un lugar, de modo que la sombra no '
                'compensa no tener a dónde caminar, como en las medidas de '
                'accesibilidad en que el estrés térmico acorta la distancia '
                'que se camina.'
            ),
            'en': (
                'Walkability sums the z-scores of access to daily living (a '
                'fresh food market, a convenience store and public transport, '
                'within 300 m), and population density and intersection '
                'density within a 1 km walk (Frank et al. 2010), taken across '
                'all sample points. By default it is attenuated by daytime '
                'thermal comfort: its percentile rank is multiplied by '
                '1 − λ·h, where h is mean daytime UTCI re-scaled from 0 at '
                '{lowValue} to 1 at {highValue}, and λ = {attenuation}. Heat '
                'reduces how walkable a place is, so '
                'shade cannot make up for having nothing to walk to, after '
                'measures of accessibility in which heat stress shortens the '
                'distance people walk.'
            ),
        },
        'variants': {
            'es': (
                'La caminabilidad puede mostrarse sin la atenuación por '
                'confort térmico. En ese caso el confort térmico entra en '
                'cambio como un indicador del entorno medioambiental, para que '
                'cuente una vez y sólo una: el índice se calcula de cada forma, '
                'cada una con sus propias metas.'
            ),
            'en': (
                'Walkability can be shown without thermal comfort '
                'attenuation. Thermal comfort then enters instead as an '
                'ambient environment indicator, so that it counts once and '
                'only once: the index is calculated each way, each against '
                'its own goalposts.'
            ),
        },
        # The information box beside the walkability setting.
        'attenuation': {
            'es': (
                'Por defecto, la caminabilidad a 300 m se atenúa por el '
                'confort térmico diurno. Su rango percentil (0 a 1) se '
                'multiplica por 1 − λ·h: h es el UTCI diurno medio reescalado '
                'entre 0 (a {lowValue}) y 1 (a {highValue}), y λ = {attenuation}, de '
                'modo que el lugar más caluroso conservaría la mitad de su '
                'caminabilidad. {observed} Sin atenuación, la caminabilidad se usa tal '
                'cual y el confort térmico se puntúa en cambio como indicador '
                'del entorno medioambiental.'
            ),
            'en': (
                'By default, walkability at 300 m is attenuated by daytime '
                'thermal comfort. Its percentile rank (0 to 1) is multiplied '
                'by 1 − λ·h: h is the mean daytime UTCI re-scaled from 0 at '
                '{lowValue} to 1 at {highValue}, and λ = {attenuation}, so a '
                'place at the hottest would keep half its walkability. '
                '{observed} Without attenuation, '
                'walkability is used as it is, and thermal comfort is scored '
                'instead as an ambient environment indicator.'
            ),
        },
        'attenuation_observed': {
            'es': (
                'En el día modelado, la mayor parte de la ciudad está entre '
                '{obsLow} y {obsHigh} (percentiles {low} y {high}, la banda '
                'sombreada): toda ella muy calurosa, así que la caminabilidad '
                'se reduce en todas partes, algo más donde hay menos sombra.'
            ),
            'en': (
                'On the modelled day most of the city lies between {obsLow} '
                'and {obsHigh} (its {low}th and {high}th percentiles, the '
                'shaded band): all of it very hot, so walkability is reduced '
                'everywhere, a little more where there is less shade.'
            ),
        },
        'attenuation_plot': {
            'es': (
                'Cada línea sigue un lugar con la caminabilidad indicada a '
                'medida que el día se vuelve más caluroso.'
            ),
            'en': (
                'Each line follows a place of the walkability shown as its '
                'daytime heat increases.'
            ),
        },
        'weights': {
            'es': (
                'El índice premia, por diseño, un perfil equilibrado: todos '
                'los dominios pesan lo mismo. La configuración permite '
                'asignar otros pesos para explorar lo que importa a cada '
                'persona; esto expresa una preferencia personal, no la '
                'evidencia de salud pública. Con pesos personalizados, las '
                'puntuaciones se recalculan en el navegador a partir de las '
                'puntuaciones promedio de cada área, por lo que son una '
                'aproximación exploratoria y no coinciden exactamente con las '
                'publicadas, que se calculan en cada punto de muestra.'
            ),
            'en': (
                'The index rewards a balanced profile by design: every domain '
                'carries the same weight. The settings allow other weights, '
                'to explore what matters to each person; this expresses '
                'personal preference, not public health evidence. With '
                'custom weights, scores are recalculated in the browser from '
                'each area\'s average scores, so they are an exploratory '
                'approximation and do not exactly match the published ones, '
                'which are calculated at each sample point.'
            ),
        },
        'utci': {
            'es': (
                'El confort térmico diurno (UTCI) se modeló para el día más '
                'caluroso '
                'de 2023 sobre superficies peatonales. La temperatura del '
                'aire, la humedad y el viento son uniformes en toda la '
                'ciudad, así que sólo varía la radiación (sol y sombra); la '
                'mayoría de los valores superan el rango en que se ajustó el '
                'polinomio del UTCI. Compare lugares cercanos entre sí, no '
                'valores absolutos ni barrios distantes.'
            ),
            'en': (
                'Daytime thermal comfort (UTCI) was modelled for the hottest '
                'day of '
                '2023 over pedestrian surfaces. Air temperature, humidity and '
                'wind are uniform across the city, so only radiation (sun and '
                'shade) varies; most values exceed the range the UTCI '
                'polynomial was fitted on. Compare nearby places with each '
                'other, not absolute values or distant neighbourhoods.'
            ),
        },
        'provisional': {
            'es': (
                'El índice es provisional: los indicadores de cada dominio y '
                'la forma de calcularlo siguen en revisión.'
            ),
            'en': (
                'The index is provisional: the indicators in each domain and '
                'how it is calculated are still under review.'
            ),
        },
        'further_reading': {
            'es': 'Lecturas adicionales',
            'en': 'Further reading',
        },
    },
}


# Domains for destination layers that are not DENUE points_of_interest entries
# (which declare their own ``domain``).
# Domains for the global default destination categories, which are declared in
# the accessibility config rather than in points_of_interest and so carry no
# domain of their own.
CATEGORY_DOMAINS = {
    'food': 'Daily essential services',
    'pos': 'Open space',
    'blue_space': 'Open space',
    'public_open_space_with_water': 'Open space',
    'pt': 'Transport',
}
LAYER_DOMAINS = {
    'aos_public_large_nodes_30m_line': 'Open space',
    'aos_public_any_nodes_30m_line': 'Open space',
    'aos_public_water_nodes_30m_line': 'Open space',
    'blue_space_nodes_30m_line': 'Open space',
    'lpugs_nodes_30m_line': 'Open space',
    'pt_stops_headway': 'Transport',
}
DEFAULT_DOMAIN = 'Other'

_QUOTED = re.compile(r"'([^']*)'")


def destination_names(spec):
    """The ``dest_name`` values a spec selects from the destinations layer."""
    if spec.get('layer') != 'destinations':
        return []
    return _QUOTED.findall(spec.get('where') or '')


def spec_domain(r, spec):
    """The reporting domain a destination spec belongs to."""
    poi = (r.config or {}).get('points_of_interest') or {}
    for name in destination_names(spec):
        domain = (poi.get(name) or {}).get('domain')
        if domain:
            return domain
    domain = spec.get('domain') or CATEGORY_DOMAINS.get(spec.get('category'))
    if domain:
        return domain
    return LAYER_DOMAINS.get(spec.get('layer'), DEFAULT_DOMAIN)


def spec_overlay(spec):
    """Where the viewer finds the points this indicator was measured to."""
    names = destination_names(spec)
    if names:
        return {'layer': 'destinations', 'dest_names': names}
    return {'layer': spec.get('layer')}


def humanise(name):
    return name.replace('_', ' ').strip().capitalize()


def _mode_configs(r):
    """Effective pedestrian and cycling accessibility configurations."""
    from _cycling_lts_network import cycling_config
    from _pedestrian_accessibility import pedestrian_config, resolve_thresholds

    shared = accessibility_config(r)
    pedestrian = pedestrian_config(r)
    ped = None
    if pedestrian is not None:
        ped = {
            'config': pedestrian,
            'thresholds': list(resolve_thresholds(pedestrian)),
            'specs': list(
                pedestrian.get('destinations') or DEFAULT_DESTINATIONS,
            ),
        }
    cycling = cycling_config(r)
    cyc = None
    if cycling is not None:
        merged = effective_config(shared, cycling)
        cyc = {
            'config': merged,
            'thresholds': sorted(
                {
                    int(d)
                    for d in (
                        merged.get('distances') or (500, 1000, 2000, 5000)
                    )
                },
            ),
            'specs': list(merged.get('destinations') or DEFAULT_DESTINATIONS),
            'networks': [k for k in CYCLE_NETWORKS],
        }
    return ped, cyc


def censored_distances(r, available):
    """The distance columns whose search stopped short, and where.

    The nearest destination is only looked for as far as the largest distance
    band an analysis measures (or its ``search_distance``): beyond it an
    area's distance is missing, not unknown, and a map or chart that reads it
    as "no data" says the wrong thing.  Returns ``{column: {distance, bands,
    access}}``: the distance searched to, the bands (for classing the
    distances by what was measured), and the access column at the largest band,
    which is 0 where nothing was found within it.
    """
    try:
        ped, cyc = _mode_configs(r)
    except Exception as e:  # reported; the dashboard exports without it
        print(f'  ! censoring distances unavailable: {e}')
        return {}
    out = {}
    for resolved, distance_prefix, access_prefix in (
        (ped, 'avg_walk_dist_', 'pct_access_walk_'),
        (cyc, 'avg_cycle_dist_', 'pct_access_cycle_'),
    ):
        if not resolved:
            continue
        bands = sorted({int(d) for d in resolved['thresholds']})
        if distance_prefix == 'avg_walk_dist_':
            from _pedestrian_accessibility import resolve_search_distance

            search = resolve_search_distance(resolved['config'], bands)
        else:
            search = bands[-1]
        classes = [d for d in bands if d <= search]
        if search > classes[-1]:
            classes.append(search)
        for column in sorted(available):
            if not column.startswith(distance_prefix):
                continue
            rest = column[len(distance_prefix) :]
            access = f'{access_prefix}{rest}_{bands[-1]}m'
            out[column] = {
                'distance': search,
                'bands': classes,
                'access': access if access in available else None,
            }
    return out


def _banded_measure(available, name, networks, thresholds, builder, meta_key):
    """A banded measure as ``{network: {distance: column}}``."""
    out = {}
    for network in networks:
        bands = {
            str(d): builder(network, name, d)
            for d in thresholds
            if builder(network, name, d) in available
        }
        if bands:
            out[network] = bands
    if not out:
        return None
    entry = dict(MEASURE_META[meta_key])
    entry['networks'] = out
    return entry


def destination_families(r, ped, cyc, available):
    """One family per destination spec, merging its walking and cycling measures."""
    families = {}
    for mode, resolved, networks in (
        ('walk', ped, [WALK]),
        ('cycle', cyc, (cyc or {}).get('networks', [])),
    ):
        if resolved is None:
            continue
        thresholds = spec_thresholds(resolved['specs'], resolved['thresholds'])
        for spec in resolved['specs']:
            name = spec['name']
            bands = list(thresholds.get(name, resolved['thresholds']))
            family = families.setdefault(
                name,
                {
                    'id': name,
                    'group': 'destination',
                    'domain': spec_domain(r, spec),
                    'category': spec.get('category'),
                    'variant': spec.get('variant'),
                    'label': {'en': humanise(name)},
                    'overlay': spec_overlay(spec),
                    'measures': {},
                },
            )
            avoid = str(spec.get('direction') or '').lower() == AVOID
            key = 'beyond' if avoid else 'access'
            builder = (
                (lambda n, s, d: beyond_column(s, d))
                if avoid
                else access_column
            )
            measure = _banded_measure(
                available,
                name,
                networks,
                bands,
                builder,
                key,
            )
            if measure:
                existing = family['measures'].get(key)
                if existing:
                    existing['networks'].update(measure['networks'])
                else:
                    family['measures'][key] = measure
            for network in networks:
                column = distance_column(network, name)
                if column not in available:
                    continue
                meta_key = 'dmextra' if network == 'dmgap' else 'distance'
                entry = family['measures'].setdefault(
                    meta_key,
                    dict(MEASURE_META[meta_key], networks={}),
                )
                entry['networks'][network] = column
    # the dismount-dependence contrast is a network of the access measure, but
    # its polarity is the opposite of the access it is derived from
    for family in families.values():
        access = family['measures'].get('access')
        if access and 'dmgap' in access.get('networks', {}):
            access.setdefault('network_direction', {})['dmgap'] = LOWER
    return [f for f in families.values() if f['measures']]


def _derived_families(
    resolved,
    networks,
    available,
    names,
    group,
    domain,
    overlay_for=None,
    labels=None,
):
    """Families for derived destinations (activity centres, composites).

    These are measured exactly like any other destination -- banded access and
    mean distance -- so they differ from ``destination_families`` only in where
    their names come from and that they have no configured spec.
    """
    families = {}
    for name in names:
        family = {
            'id': name,
            'group': group,
            'domain': domain,
            'label': (labels or {}).get(name) or {'en': humanise(name)},
            'measures': {},
        }
        if overlay_for:
            overlay = overlay_for(name)
            if overlay:
                family['overlay'] = overlay
        measure = _banded_measure(
            available,
            name,
            networks,
            resolved['thresholds'],
            access_column,
            'access',
        )
        if measure:
            family['measures']['access'] = measure
        distances = {
            network: distance_column(network, name)
            for network in networks
            if distance_column(network, name) in available
        }
        if distances:
            family['measures']['distance'] = dict(
                MEASURE_META['distance'],
                networks=distances,
            )
        if family['measures']:
            families[name] = family
    return families


def _merge_families(into, more):
    for name, family in more.items():
        existing = into.get(name)
        if not existing:
            into[name] = family
            continue
        for key, measure in family['measures'].items():
            current = existing['measures'].get(key)
            if current:
                current['networks'].update(measure['networks'])
            else:
                existing['measures'][key] = measure


def activity_centre_names(config):
    """The derived activity-centre destination names, per definition and tier."""
    names, labels = [], {}
    for set_name, definition in activity_centre_definitions(config).items():
        for tier in definition['tiers']:
            name = (
                f'activity_centre_{tier}'
                if set_name == STANDARD_SET
                else f'activity_centre_{set_name}_{tier}'
            )
            names.append(name)
            categories = ', '.join(definition['categories'])
            labels[name] = {
                'en': (
                    f'{humanise(set_name)} activity centre ({tier}), '
                    f'{definition["walk_threshold"]} m: {categories}'
                ),
            }
    return names, labels


def combined_access_names(config, specs):
    """The composite 'everything reachable' destination names."""
    names, labels = [], {}
    for set_name, categories in combined_access_sets(config, specs).items():
        for variant in ('strict', 'lenient'):
            name = (
                f'all_{variant}'
                if set_name == STANDARD_SET
                else f'all_{set_name}_{variant}'
            )
            names.append(name)
            labels[name] = {
                'en': (f'All of {", ".join(categories)} ({variant} variant)'),
            }
    return names, labels


def diversity_families(r, ped, available):
    """Diversity, richness and per-sub-type count families."""
    if ped is None:
        return []
    sets = diversity_sets(ped['config'])
    families = []
    for set_name, spec in sets.items():
        bands = sorted(
            {int(d) for d in (spec['distances'] or ped['thresholds'])},
        )
        family = {
            'id': f'diversity_{set_name}',
            'group': 'diversity',
            'domain': 'Diversity',
            'label': {'en': f'{humanise(set_name)} diversity'},
            'measures': {},
            'groups': list(spec['groups']),
        }
        for key, builder in (
            ('diversity', diversity_column),
            ('richness', richness_column),
        ):
            columns = {
                str(d): builder(set_name, d)
                for d in bands
                if builder(set_name, d) in available
            }
            if columns:
                family['measures'][key] = dict(
                    MEASURE_META[key],
                    networks={WALK: columns},
                )
        for group in spec['groups']:
            columns = {
                str(d): count_column(set_name, group, d)
                for d in bands
                if count_column(set_name, group, d) in available
            }
            if columns:
                family['measures'].setdefault(
                    'count',
                    dict(MEASURE_META['count'], groups={}),
                )['groups'][group] = {WALK: columns}
        if family['measures']:
            families.append(family)
    return families


NEUTRAL = 'neutral'

# Standalone indicators: one column each, grouped for the picker.  Direction is
# stated per variable because it is not derivable from the name -- a higher
# albedo is better and a higher land surface temperature is worse, and both are
# urban_heat_ columns.
URBAN_HEAT_DIRECTION = {
    'urban_heat_land_surface_albedo': HIGHER,
    'urban_heat_ndvi': HIGHER,
    'urban_heat_subnational_hdi': HIGHER,
    'urban_heat_adaptive_capability_index': HIGHER,
}
STANDALONE_GROUPS = [
    {
        'id': 'walkability',
        'domain': 'Walkability',
        'label': {'es': 'Caminabilidad', 'en': 'Walkability'},
        'columns': [
            'local_walkability',
            'local_daily_living',
            'local_nh_population_density',
            'local_nh_intersection_density',
        ],
        'direction': HIGHER,
    },
    {
        'id': 'urban_heat',
        'domain': 'Urban heat',
        'label': {
            'es': 'Vulnerabilidad al calor urbano',
            'en': 'Urban heat vulnerability',
        },
        'columns': (
            ['urban_heat_guhvi', 'urban_heat_guhvi_class']
            + [
                f'urban_heat_{parent}'
                for parent, _ in URBAN_HEAT_STRUCTURE_PARENTS
            ]
            + [
                f'urban_heat_{sub}'
                for _, subs in URBAN_HEAT_STRUCTURE_PARENTS
                for sub in subs
            ]
            + ['pct_urban_heat_guhvi_class_5_most_vulnerable']
        ),
        'direction': LOWER,
        'direction_overrides': URBAN_HEAT_DIRECTION,
    },
    {
        'id': 'context',
        'domain': 'Context',
        'label': {'es': 'Contexto', 'en': 'Context'},
        'columns': [
            'pop_est',
            'pop_per_sqkm',
            'area_sqkm',
            'intersection_count',
            'intersections_per_sqkm',
        ],
        'direction': NEUTRAL,
    },
]


def core_access_family(available):
    """The fixed indicators.yml 500 m access scores, kept as their own group.

    These are the globally comparable core measures; they are not part of the
    configurable banded vocabulary and always sit at 500 m.
    """
    columns = sorted(
        c
        for c in available
        if c.startswith('pct_access_500m_') and c.endswith('_score')
    )
    if not columns:
        return None
    return {
        'id': 'core_access',
        'group': 'standalone',
        'domain': 'Core access (500 m)',
        # distinct from the walking analysis's own 500 m band: these are the
        # fixed, globally comparable measures, over OpenStreetMap destinations
        'label': {
            'es': 'Comparación global: acceso a 500 m (OSM)',
            'en': 'Global comparison: access within 500 m (OSM)',
        },
        'direction': HIGHER,
        'measures': {
            'value': dict(
                MEASURE_META['value'],
                direction=HIGHER,
                units='percent',
                variables={c: {WALK: c} for c in columns},
            ),
        },
    }


def standalone_families(available):
    families = []
    for group in STANDALONE_GROUPS:
        columns = [c for c in group['columns'] if c in available]
        if not columns:
            continue
        overrides = group.get('direction_overrides') or {}
        families.append(
            {
                'id': group['id'],
                'group': 'standalone',
                'domain': group['domain'],
                'label': group['label'],
                'direction': group['direction'],
                'measures': {
                    'value': dict(
                        MEASURE_META['value'],
                        direction=group['direction'],
                        variables={c: {WALK: c} for c in columns},
                        variable_direction={
                            c: overrides[c] for c in columns if c in overrides
                        },
                    ),
                },
            },
        )
    return families


def _normalise(text):
    """Fold a label to a comparison key: case, accents and spacing all vary."""
    import unicodedata

    folded = unicodedata.normalize('NFKD', str(text or ''))
    folded = ''.join(c for c in folded if not unicodedata.combining(c))
    folded = re.sub(r'[^a-z0-9]+', ' ', folded.lower())
    return folded.strip()


# Spellings in the source material that must not become a second theme: the
# intervention sheet writes one theme both with and without a space before its
# bracket, and spells "transporte" as "transprote" in one place.
_SPELLING = {'transprote': 'transporte'}


def _theme_key(text):
    return ' '.join(_SPELLING.get(w, w) for w in _normalise(text).split())


def resolve_themes(config):
    """The workshop's themes: ``{id: {label, color, families, keys}}``.

    Themes are the vocabulary the workshop itself uses -- its printed materials
    are organised and coloured by them -- so they are declared in the region
    config rather than derived from the analysis, which knows only its own
    reporting domains.
    """
    themes = {}
    for key, spec in (config.get('themes') or {}).items():
        spec = spec or {}
        label = _labels(spec.get('label'), key)
        matches = spec.get('match') or [label.get('es'), label.get('en'), key]
        themes[key] = {
            'id': key,
            'label': label,
            'color': spec.get('color'),
            'families': list(spec.get('families') or []),
            # every spelling this theme should be recognised by when the
            # intervention sheet is matched against it
            'keys': sorted({_theme_key(m) for m in matches if m}),
        }
    return themes


def apply_labels(config, families, descriptions):
    """Overlay configured ``{es, en}`` labels onto the derived ones.

    A family name is otherwise a humanised column name -- "Denue pharmacy" --
    which is neither Spanish nor what anyone calls the thing.
    """
    labels = config.get('labels') or {}
    family_labels = labels.get('families') or {}
    variable_labels = labels.get('variables') or {}
    missing = []
    for family in families:
        configured = family_labels.get(family['id'])
        if configured:
            family['label'] = _labels(configured, family['id'])
        elif 'es' not in (family.get('label') or {}):
            missing.append(family['id'])
    for variable, configured in variable_labels.items():
        if variable in descriptions:
            descriptions[variable]['label'] = _labels(configured, variable)
    if missing:
        print(
            f'  ! {len(missing)} families have no Spanish label: '
            f'{", ".join(missing[:6])}' + (' ...' if len(missing) > 6 else ''),
        )
    return families


def assign_themes(themes, families):
    """Tag each family with its theme, and each theme with its families.

    An unthemed family is reported rather than silently dropped: with themes as
    the primary grouping, an indicator that cannot be navigated to may as well
    not have been produced.
    """
    by_family = {}
    for theme in themes.values():
        for family_id in theme['families']:
            by_family[family_id] = theme['id']
    unthemed = []
    for family in families:
        theme = by_family.get(family['id'])
        if theme:
            family['theme'] = theme
        else:
            unthemed.append(family['id'])
    if unthemed:
        print(
            f'  ! {len(unthemed)} families have no theme: '
            f'{", ".join(unthemed[:6])}'
            + (' ...' if len(unthemed) > 6 else ''),
        )
    known = {f['id'] for f in families}
    for theme in themes.values():
        unknown = [f for f in theme['families'] if f not in known]
        if unknown:
            print(
                f'  ! theme "{theme["id"]}" lists families not produced here: '
                f'{", ".join(unknown)}',
            )
        theme['families'] = [f for f in theme['families'] if f in known]
    return families


def load_interventions(path, themes):
    """Interventions and their impacts, from the workshop's matching sheet.

    Columns: TEMA | number | intervention | core measure | Reimagina | impact.
    TEMA and the intervention columns are merged down in the source, so they
    are carried forward; an intervention appears once per core measure it
    touches and is folded back into a single entry here.
    """
    if not path:
        return []
    resolved = path
    if not os.path.isabs(resolved):
        resolved = _data_path(path)
    if not os.path.exists(resolved):
        print(f'  ! interventions sheet not found, skipping: {resolved}')
        return []
    import openpyxl

    workbook = openpyxl.load_workbook(resolved, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    by_key = {
        k: theme['id'] for theme in themes.values() for k in theme['keys']
    }
    unmatched = set()
    entries = {}
    carried = ['', '', '', '']
    splitter = re.compile(r':|\s{2,}')
    for row in rows[1:]:
        cells = [
            '' if c is None else str(c).replace('\n', ' ').strip() for c in row
        ]
        cells += [''] * (6 - len(cells))
        tema, number, name, measure, reimagina, impact = cells[:6]
        for i, value in enumerate((tema, number, name, reimagina)):
            if value:
                carried[i] = value
        tema, number, name, reimagina = carried
        if not number or not name:
            continue
        parts = splitter.split(name, maxsplit=1)
        entry = entries.setdefault(
            number,
            {
                'n': number,
                # the sheet packs a title and its description into one cell,
                # separated by a colon or a run of spaces
                'name': parts[0].strip(),
                'description': parts[1].strip() if len(parts) > 1 else '',
                'reimagina': reimagina,
                'themes': {},
                'core_measures': [],
            },
        )
        theme_id = by_key.get(_theme_key(tema))
        if not theme_id and tema:
            unmatched.add(tema)
        if measure and measure not in entry['core_measures']:
            entry['core_measures'].append(measure)
        # Impacts are recorded per theme, not pooled.  An intervention is only
        # offered for a theme where the sheet actually states what it would do
        # there: "Ciclovía" is tagged to six themes but carries no impact under
        # Essential services, and showing it against fresh food -- as the only
        # entry -- read as noise.
        if theme_id and impact:
            impacts = entry['themes'].setdefault(theme_id, [])
            if impact not in impacts:
                impacts.append(impact)
    if unmatched:
        print(
            '  ! intervention themes not matched to a configured theme: '
            f'{"; ".join(sorted(unmatched))}',
        )
    kept = [e for e in entries.values() if e['themes']]
    dropped = len(entries) - len(kept)
    print(
        f'  Interventions: {len(kept)} from {os.path.basename(resolved)}'
        + (f' ({dropped} with no stated impact, omitted)' if dropped else ''),
    )
    return kept


def load_crosswalk(path):
    """ULI core-measure crosswalk as ``{variable: {...}}``, or {}.

    The crosswalk states which project measure each GHSCI variable answers and
    with what relation and direction.  It is optional: without it the dashboard
    still works, it just has no core-measure view.
    """
    if not path:
        return {}
    resolved = path
    if not os.path.isabs(resolved):
        resolved = _data_path(path)
    if not os.path.exists(resolved):
        print(f'  ! crosswalk not found, skipping: {resolved}')
        return {}
    import yaml

    with open(resolved, encoding='utf-8') as f:
        data = yaml.safe_load(f) or {}
    out = {}
    for measure, entry in (data.get('core_measures') or {}).items():
        for sub in entry.get('sub_variables') or []:
            for variable in sub.get('variables') or []:
                out[variable] = {
                    'measure': measure,
                    'core_measure': entry.get('core_measure'),
                    'domains': entry.get('domains'),
                    'sub_variable': sub.get('name'),
                    'relation': sub.get('relation'),
                    'direction': sub.get('direction'),
                }
    print(f'  Crosswalk: {len(out)} variables mapped to core measures')
    return out


# How each measure reads in Spanish, given the indicator, the network and the
# band.  Composed rather than translated: every part is already in Spanish, so
# this gives a real Spanish description for the ~310 configurable columns
# without inventing prose for any of them.
SPANISH_DESCRIPTION = {
    'access': lambda f, net, band: (
        f'Proporción de puntos de muestreo con acceso a {band} m de: {f} '
        f'— red: {net}'
    ),
    'beyond': lambda f, net, band: (
        f'Proporción de puntos de muestreo a más de {band} m de: {f} '
        f'— red: {net}'
    ),
    'distance': lambda f, net, band: (
        f'Distancia media por la red hasta lo más cercano: {f} — red: {net}'
    ),
    'dmextra': lambda f, net, band: (
        f'Metros adicionales de pedaleo para evitar desmontar: {f}'
    ),
    'count': lambda f, net, band: (
        f'Número medio alcanzable a {band} m: {f} — red: {net}'
    ),
    # the family label already begins "Diversidad de ..." / "Riqueza de ...",
    # so these say the band and the network and let it speak for itself
    'diversity': lambda f, net, band: f'{f}, a {band} m — red: {net}',
    'richness': lambda f, net, band: f'{f}, a {band} m — red: {net}',
}


def compose_spanish_descriptions(families, descriptions):
    """Give every column a Spanish description built from translated parts.

    The generated data dictionary is English only, and translating its ~1,000
    sentences is not something to invent.  But the columns this dashboard
    exposes are almost all of the form "measure X of destination Y over network
    Z at band B", and every one of those parts already has a Spanish label -- so
    the sentence can be composed rather than translated.  Standalone variables
    (urban heat, walkability, context) carry a configured Spanish label instead.
    """
    composed = 0
    for family in families:
        spanish = (family.get('label') or {}).get('es')
        if not spanish:
            continue
        for key, measure in family['measures'].items():
            write = SPANISH_DESCRIPTION.get(key)
            if not write:
                continue
            for network, value in (measure.get('networks') or {}).items():
                network_es = (
                    (NETWORKS.get(network) or {}).get('label') or {}
                ).get('es') or network
                bands = {None: value} if isinstance(value, str) else value
                for band, column in bands.items():
                    entry = descriptions.setdefault(column, {})
                    entry['es'] = write(spanish, network_es, band)
                    composed += 1
            for by_network in (measure.get('groups') or {}).values():
                for network, bands in by_network.items():
                    network_es = (
                        (NETWORKS.get(network) or {}).get('label') or {}
                    ).get('es') or network
                    for band, column in bands.items():
                        entry = descriptions.setdefault(column, {})
                        entry['es'] = write(spanish, network_es, band)
                        composed += 1
    # standalone variables: the configured label is the description
    for column, entry in descriptions.items():
        if 'es' not in entry and entry.get('label', {}).get('es'):
            entry['es'] = entry['label']['es']
            composed += 1
    print(f'  Spanish descriptions composed for {composed} columns')
    return descriptions


_CODES = re.compile(r'codigo_act\s+IN\s*\(([^)]*)\)', re.IGNORECASE)


def destination_provenance(r, spec):
    """Where the points an indicator measured to actually came from.

    Read from the region's own ``points_of_interest`` entries, including the
    SCIAN activity codes the DENUE filter selects -- which is the level of
    detail somebody checking a Mexicali result will want, and which is
    otherwise buried in an ogr2ogr ``-where`` clause.
    """
    poi = (r.config or {}).get('points_of_interest') or {}
    sources = []
    for name in destination_names(spec):
        entry = poi.get(name)
        if not entry:
            continue
        codes = _CODES.search(str(entry.get('data') or ''))
        sources.append(
            {
                'dest_name': name,
                'name': entry.get('dest_name_full') or name,
                'codes': (
                    [c.strip().strip("'\"") for c in codes.group(1).split(',')]
                    if codes
                    else []
                ),
                'source': entry.get('source') or '',
                'publication_date': str(entry.get('publication_date') or ''),
                'url': entry.get('url') or '',
                'licence': entry.get('licence') or '',
                'citation': entry.get('citation') or '',
            },
        )
    return sources


def analysis_rules(r):
    """The rules the analysis applied, for the info panel's method section.

    Emitted once rather than per family: the pedestrian filter and the traffic
    stress thresholds are properties of the region's analysis, not of any one
    indicator.
    """
    from _cycling_lts_network import ADT_BY_GROUP, LTS_IMPED, cycling_config

    network = (r.config or {}).get('network') or {}
    # 'openstreetmap_query' is the current name for the OSMnx custom filter.
    # 'network' (active travel branch) and 'pedestrian' are its earlier
    # names, still read so this exporter works whichever way a region was
    # configured.  The emitted key is left as 'pedestrian_filter', which the
    # site consumes.
    openstreetmap_query = str(
        network.get('openstreetmap_query')
        or network.get('network')
        or network.get('pedestrian')
        or '',
    ).strip()
    rules = {
        'pedestrian_filter': openstreetmap_query,
        'retains_private_access': 'access' in openstreetmap_query,
        'osmnx_retain_all': bool(network.get('osmnx_retain_all')),
    }
    cycling = cycling_config(r)
    if cycling is not None:
        rules['cycling'] = {
            'adt_by_group': dict(ADT_BY_GROUP),
            'lts_impedance': {str(k): v for k, v in LTS_IMPED.items()},
            'danger_weight': (cycling or {}).get('danger_weight'),
            'dismount_weight': (cycling or {}).get('dismount_weight'),
        }
    return rules


def _family_columns(family):
    """Every physical column any measure of a family refers to."""
    columns = set(family.get('extra_columns') or [])
    for measure in family['measures'].values():
        for holder in ('networks', 'variables'):
            for value in (measure.get(holder) or {}).values():
                if isinstance(value, dict):
                    columns.update(value.values())
                else:
                    columns.add(value)
        for by_network in (measure.get('groups') or {}).values():
            for bands in by_network.values():
                columns.update(bands.values())
    return columns


def _standalone_family(
    family_id,
    domain,
    label,
    columns,
    direction,
    overrides=None,
):
    """A standalone family of single-column variables."""
    overrides = overrides or {}
    return {
        'id': family_id,
        'group': 'standalone',
        'domain': domain,
        'label': label,
        'direction': direction,
        'measures': {
            'value': dict(
                MEASURE_META['value'],
                direction=direction,
                variables={c: {WALK: c} for c in columns},
                variable_direction={
                    c: overrides[c] for c in columns if c in overrides
                },
            ),
        },
    }


CATCHMENT_PREFIXES = (
    'pct_access_euclid_',
    'pct_beyond_euclid_',
    'avg_euclid_dist_',
)


def catchment_families(available):
    """Straight-line catchment indicators (see _euclidean_accessibility)."""
    columns = sorted(c for c in available if c.startswith(CATCHMENT_PREFIXES))
    if not columns:
        return []
    return [
        _standalone_family(
            'catchments',
            'Catchments',
            {
                'es': 'Cobertura en línea recta',
                'en': 'Straight-line catchments',
            },
            columns,
            HIGHER,
            {c: LOWER for c in columns if c.startswith('avg_euclid_dist_')},
        ),
    ]


def linkage_families(r, available):
    """Externally prepared indicators, one family per linked source.

    Their labels, units, descriptions and directions are the ones configured
    with them (see _linkage_indicators), carried as ``family['linked']`` for
    :func:`apply_linked_metadata`, since the data dictionary cannot know them.
    """
    try:
        from _linkage_indicators import linkage_config

        specs = linkage_config(r) or {}
    except ValueError as e:
        print(f'  Linked indicators omitted from the dashboard: {e}')
        return []
    families = []
    for name, spec in specs.items():
        meta = {
            spec['outputs'][column]: values
            for column, values in spec['columns'].items()
        }
        columns = [c for c in meta if c in available]
        if not columns:
            continue
        family = _standalone_family(
            f'linked_{name}',
            'Linked indicators',
            _labels(spec.get('label'), humanise(name)),
            columns,
            HIGHER,
            {c: LOWER for c in columns if meta[c].get('direction') == LOWER},
        )
        family['linked'] = {c: meta[c] for c in columns}
        family['sources'] = [
            {
                'dest_name': name,
                'name': _labels(spec.get('label'), name).get('en', name),
                'codes': [],
                'source': spec.get('source') or '',
                'publication_date': '',
                'url': '',
                'licence': spec.get('licence') or '',
                'citation': spec.get('citation') or '',
            },
        ]
        if spec.get('notes'):
            family['note'] = _labels(spec['notes'], '')
        families.append(family)
    return families


def apply_linked_metadata(families, descriptions):
    """Describe linked columns as configured, which describe_variable cannot."""
    for family in families:
        for column, meta in (family.get('linked') or {}).items():
            entry = descriptions.get(column)
            if entry is None:
                continue
            entry['category'] = 'Linked indicators'
            if meta.get('description'):
                entry['en'] = meta['description']
            if meta.get('units'):
                entry['units'] = str(meta['units'])
            if meta.get('direction'):
                entry['direction'] = meta['direction']
            if meta.get('label'):
                entry.setdefault('label', _labels(meta['label'], column))
    return descriptions


def walkability_variant_families(r, available):
    """Walkability at each configured distance, and its heat variants.

    Labelled from the configuration (distance, heat measures and form), as
    names like ``walk_idx_300_gtm`` say nothing to a reader.
    """
    try:
        from _walkability_variants import (
            DAILY_LIVING_PREFIX,
            variants,
            walkability_config,
        )

        config = walkability_config(r)
    except ValueError as e:
        print(f'  Walkability variants omitted from the dashboard: {e}')
        return []
    if config is None:
        return []
    heat = {
        letter: _labels(m.get('label'), m['variable'])
        for letter, m in config['heat'].items()
    }
    forms = {
        'additive': {'es': 'aditiva', 'en': 'additive'},
        'multiplicative': {'es': 'atenuación', 'en': 'attenuation'},
    }
    labels, units, columns = {}, {}, []
    for variant in variants(config):
        column = variant['column'][len('sp_') :]
        if column not in available:
            continue
        columns.append(column)
        # a sum of z-scores, or (attenuated) a percentile rank from 0 to 1
        units[column] = (
            'index 0-1'
            if variant['form'] == 'multiplicative'
            else 'index (sum of z-scores)'
        )
        label = {
            'es': f"Caminabilidad ({variant['distance']} m)",
            'en': f"Walkability ({variant['distance']} m)",
        }
        if variant['heat']:
            for lang in ('es', 'en'):
                joined = (' y ' if lang == 'es' else ' and ').join(
                    heat[c].get(lang, c) for c in variant['heat']
                )
                label[lang] += (
                    f" {'+' if variant['form'] == 'additive' else '×'} "
                    f"{joined} ({forms[variant['form']][lang]})"
                )
        labels[column] = label
    for distance in config['distances']:
        column = f'{DAILY_LIVING_PREFIX[len("sp_") :]}{distance}'
        if column in available:
            columns.append(column)
            units[column] = 'destinations'
            labels[column] = {
                'es': f'Puntaje de vida diaria ({distance} m)',
                'en': f'Daily living score ({distance} m)',
            }
    if not columns:
        return []
    family = _standalone_family(
        'walkability_variants',
        'Walkability',
        {
            'es': 'Caminabilidad (300 y 500 m, y ajustada por calor)',
            'en': 'Walkability (300 and 500 m, and heat adjusted)',
        },
        columns,
        HIGHER,
    )
    family['variable_labels'] = labels
    family['variable_units'] = units
    return [family]


_BAND_SUFFIX = re.compile(r'_(\d+)m$')

# the parts of an index that are not scores, named in both languages
COMPOSITE_PARTS = {
    'mean': {
        'es': 'nivel medio, antes de la penalización',
        'en': 'mean level, before the penalty',
    },
    'penalty': {
        'es': 'penalización por desequilibrio',
        'en': 'imbalance penalty',
    },
}


def _composite_parameters(r):
    """The parameters each composite index was last scored with, if recorded."""
    try:
        from _composite_index import recorded_parameters

        return recorded_parameters(r) or {}
    except Exception as e:
        print(f'  Composite index parameters not read: {e}')
        return {}


def composite_label(labels, indicator):
    """An ``{es, en}`` label for one indicator of a composite index.

    A label configured on the indicator wins.  Otherwise the dashboard's own
    label for what the indicator measures -- the destination family, or the
    standalone variable -- which is already in both languages and is the name a
    reader has already met elsewhere in the dashboard.
    """
    from _composite_index import short_name

    if indicator.get('label'):
        return _labels(indicator['label'], indicator['id'])
    key = short_name(indicator['variable'])
    band = _BAND_SUFFIX.search(key)
    base = _BAND_SUFFIX.sub('', key)
    families = labels.get('families') or {}
    variables = labels.get('variables') or {}
    found = (
        families.get(base)
        or variables.get(indicator['variable'])
        or variables.get(key)
        or next(
            (v for k, v in variables.items() if k.endswith(f'_{key}')),
            None,
        )
    )
    if not found:
        return {'en': humanise(key)}
    found = _labels(found, key)
    if band:
        found = {
            lang: f'{text} ({band.group(1)} m)' for lang, text in found.items()
        }
    return found


def _join_labels(first, second):
    """'first · second' in each language both are written in."""
    return {
        lang: f'{first[lang]} · {second[lang]}'
        for lang in ('es', 'en')
        if (first or {}).get(lang) and (second or {}).get(lang)
    }


# How a sample point variable is named once averaged up to areas, as
# _12_aggregation's propagation renames it (see calc_pedestrian_indicators,
# calc_cycling_indicators, calc_euclidean_indicators and the walkability,
# linkage and urban heat steps): (pattern, replacement), tried in order.
SAMPLE_POINT_OUTPUTS = (
    (r'^sp_walk_access_', 'pct_access_walk_'),
    (r'^sp_walk_beyond_', 'pct_beyond_walk_'),
    (r'^sp_walk_nearest_node_', 'avg_walk_dist_'),
    (r'^sp_walk_count_', 'avg_count_walk_'),
    (r'^sp_walk_diversity_', 'avg_diversity_walk_'),
    (r'^sp_walk_richness_', 'avg_richness_walk_'),
    (r'^sp_walk_idx_', 'walk_idx_'),
    (r'^sp_walk_dl_', 'walk_dl_'),
    (r'^sp_euclid_access_', 'pct_access_euclid_'),
    (r'^sp_euclid_beyond_', 'pct_beyond_euclid_'),
    (r'^sp_euclid_dist_', 'avg_euclid_dist_'),
    (r'^sp_cycle_(.*?)nearest_node_', r'avg_cycle_dist_\1'),
    (r'^sp_cycle_(.*?)access_', r'pct_access_cycle_\1'),
    # the fixed indicators.yml access scores
    (r'^sp_access_', 'pct_access_500m_'),
    (r'^sp_nearest_node_', 'avg_dist_'),
    # linked indicators and urban heat keep their name less 'sp_'
    (r'^sp_', ''),
)


def area_column_for(variable, available):
    """The area column holding a sample point variable's average, or None.

    Used to report an index indicator's raw value -- metres, a percentage, a
    temperature -- beside its normalised score.
    """
    for pattern, replacement in SAMPLE_POINT_OUTPUTS:
        if re.match(pattern, variable):
            candidate = re.sub(pattern, replacement, variable, count=1)
            if candidate in available:
                return candidate
    return None


def reading_columns(indicator, available):
    """The area columns an indicator's plain-language reading draws on.

    ``value`` is its own average (``area_column``).  For a distance scored
    against a soft threshold, ``access`` is the share of the population
    within the threshold, where that band was measured
    (``pct_access_walk_<destination>_<t>m``).  An indicator's configured
    ``reading_columns`` add to or replace these; only columns exported are
    kept.
    """
    found = {}
    if indicator.get('area_column'):
        found['value'] = indicator['area_column']
    threshold = indicator.get('soft_threshold')
    variable = str(indicator.get('variable') or '')
    if threshold:
        metres = int(round(float(threshold)))
        for pattern, access in (
            (r'^sp_walk_nearest_node_(.+)$', r'pct_access_walk_\1_{t}m'),
            (
                r'^sp_cycle_(.*?)nearest_node_(.+)$',
                r'pct_access_cycle_\1\2_{t}m',
            ),
            (r'^sp_euclid_dist_(.+)$', r'pct_access_euclid_\1_{t}m'),
        ):
            if re.match(pattern, variable):
                column = re.sub(pattern, access, variable).format(t=metres)
                if column in available:
                    found['access'] = column
                break
    found.update(indicator.get('reading_columns') or {})
    return {k: v for k, v in found.items() if v in available}


def composite_families(r, available, config=None):
    """One family per configured composite index, carrying its structure.

    Its variables are every score the index writes -- the index, then each
    domain followed by its indicators, then the mean level and penalty -- so
    any part of the index can be mapped.  ``family['composite']`` describes how
    those scores fit together, for the viewer's profile chart.
    """
    from _composite_index import (
        POSITIVE,
        composite_index_config,
        index_columns,
        index_structure,
    )

    try:
        specs = composite_index_config(r) or {}
    except ValueError as e:
        print(f'  Composite indices omitted from the dashboard: {e}')
        return []
    labels = (config or {}).get('labels') or {}
    parameters = _composite_parameters(r) if specs else {}
    families = []
    for name, spec in specs.items():
        # a variant is presented within its base index's family, as a setting
        if spec.get('variant_of'):
            continue
        columns = [
            c
            for c in index_columns(spec)
            if c in available and not c.endswith('_n')
        ]
        if not columns:
            continue
        structure = index_structure(spec, parameters.get(name))
        structure['label'] = _labels(spec.get('label'), humanise(name))
        for indicator in structure['indicators']:
            indicator['label'] = composite_label(labels, indicator)
            if indicator['column'] not in available:
                indicator['column'] = None
            # the indicator's own value, averaged over each area, reported
            # beside its normalised score
            indicator['area_column'] = area_column_for(
                indicator['variable'],
                available,
            )
            # what an area's value means in words: the configured
            # template's columns, else its average value and, for a distance
            # scored against a threshold, the share of the population within
            indicator['reading_columns'] = reading_columns(
                indicator,
                available,
            )
        by_id = {i['id']: i for i in structure['indicators']}
        for domain in structure['domains']:
            domain['label'] = _labels(
                domain['label'],
                humanise(domain['name'] or name),
            )
            if domain['column'] not in available:
                domain['column'] = None
            # a member as the index describes it, with its share of this domain
            # and the part of the domain it answers to here
            domain['indicators'] = [
                {
                    **by_id[i['id']],
                    'share': i['share'],
                    'subdomain': i.get('subdomain'),
                }
                for i in domain['indicators']
            ]
        parts = (f'_mean', f'_penalty')
        ordered = [c for c in columns if not c.endswith(parts)] + [
            c for c in columns if c.endswith(parts)
        ]
        family = _standalone_family(
            f'composite_{name}',
            'Composite indices',
            structure['label'],
            ordered,
            HIGHER if spec['phenomenon'] == POSITIVE else LOWER,
            # a larger penalty is a less balanced profile, whatever the
            # phenomenon measured
            {f'index_{name}_penalty': LOWER},
        )
        variants, extra = composite_variants(specs, spec, available)
        if variants:
            structure['variants'] = variants
        # the variants' scores and every indicator's raw value are in the
        # tiles, so the viewer can map any variant and report the values, but
        # are not offered as variables of their own: the settings choose
        # between variants
        raw = (
            [
                i['area_column']
                for i in structure['indicators']
                if i.get('area_column')
            ]
            + [
                c
                for i in structure['indicators']
                for c in (i.get('reading_columns') or {}).values()
            ]
            + [
                c
                for v in variants
                for c in (v.get('area_columns') or {}).values()
            ]
        )
        extra = list(dict.fromkeys(extra + raw))
        if extra:
            family['extra_columns'] = extra
        family['composite'] = structure
        families.append(family)
    return families


def composite_variants(specs, spec, available):
    """The variants of an index, as the dashboard's settings present them.

    Each gives the columns holding its index, mean level, penalty and domain
    scores, and the normalised scores of the indicators it swaps or activates
    (with their raw area values, ``area_columns``); its other indicator scores
    are the base index's.  The base index itself comes first.  Any further
    descriptive settings of a variant (``attenuation``, ``walk``...) are passed
    through.  Returns ``(variants, columns)``, the columns being every variant
    column exported, beyond the base index's own.
    """
    from _composite_index import (
        OUTPUT_PREFIX,
        component_column,
        effective_weights,
        indicator_score_column,
        scored_domains,
    )

    described = spec.get('variants') or []
    replaces = spec.get('variant_replaces')
    if not described:
        return [], []
    variants, extra = [], []
    for item in described:
        name = item['name']
        if name not in specs:
            continue
        base = f'{OUTPUT_PREFIX}{name}'
        columns = {
            'index': base,
            'mean': f'{base}_mean',
            'penalty': f'{base}_penalty',
        }
        domains = {
            d['name']: component_column(name, d['name'])
            for d in scored_domains(specs[name])
            if d['name'] is not None
        }
        if columns['index'] not in available:
            continue
        variables = {i['id']: i['variable'] for i in specs[name]['indicators']}
        indicators, area_columns = {}, {}
        for indicator in [replaces, *(item.get('activates') or [])]:
            column = indicator_score_column(name, indicator)
            if column in available:
                indicators[indicator] = column
            raw = area_column_for(variables.get(indicator, ''), available)
            if raw:
                area_columns[indicator] = raw
        described = {
            k: v
            for k, v in item.items()
            if k not in ('key', 'name', 'label', 'variable', 'activates')
        }
        entry = {
            **described,
            'key': item['key'],
            'name': name,
            'label': _labels(item.get('label'), humanise(name)),
            'variable': item.get('variable'),
            'activates': list(item.get('activates') or []),
            'columns': {k: v for k, v in columns.items() if v in available},
            'domains': {k: v for k, v in domains.items() if v in available},
            'indicators': indicators,
            'area_columns': area_columns,
            # what each indicator counts for in this variant, which scores
            # a different set (activating thermal comfort, say)
            'effective_weights': {
                k: round(v['effective'], 6)
                for k, v in effective_weights(specs[name]).items()
            },
        }
        variants.append(entry)
        if name != spec['name']:
            extra += list(entry['columns'].values())
            extra += list(entry['domains'].values())
            extra += list(entry['indicators'].values())
    return variants, extra


# the units of composite scores reported relative to the reference
COMPOSITE_DIFFERENCE_UNITS = (
    'points from the reference (0 = study region average)'
)


def apply_composite_labels(families, descriptions):
    """Name each composite index column for the part of the index it holds.

    Applied after the configured variable labels, which therefore still win.
    """
    for family in families:
        structure = family.get('composite')
        if not structure:
            continue
        columns = structure['columns']
        named = {
            columns['index']: structure['label'],
            columns['mean']: _join_labels(
                structure['label'],
                COMPOSITE_PARTS['mean'],
            ),
            columns['penalty']: _join_labels(
                structure['label'],
                COMPOSITE_PARTS['penalty'],
            ),
        }
        for domain in structure['domains']:
            if domain['column']:
                named[domain['column']] = domain['label']
        # scores reported as differences from the reference are points from
        # it, not an index about 100
        if structure.get('reference_value') == 0:
            scored = (
                [columns['index'], columns['mean']]
                + [d['column'] for d in structure['domains'] if d['column']]
                + [i['column'] for i in structure['indicators'] if i['column']]
            )
            for variant in structure.get('variants') or []:
                scored += [
                    c
                    for k, c in (variant.get('columns') or {}).items()
                    if k != 'penalty'
                ]
                scored += list((variant.get('domains') or {}).values())
                scored += list((variant.get('indicators') or {}).values())
            for column in scored:
                if column in descriptions:
                    described = descriptions[column]
                    described['units'] = COMPOSITE_DIFFERENCE_UNITS
                    for lang in ('en', 'es'):
                        if isinstance(described.get(lang), str):
                            described[lang] = (
                                described[lang]
                                .replace(
                                    '(100 = reference)',
                                    '(0 = the study region average)',
                                )
                                .replace(
                                    '(100 = referencia)',
                                    '(0 = promedio de la región de estudio)',
                                )
                            )
        # an indicator's score is written once, whatever domains it counts
        # towards, so it is named for the index rather than for a domain
        indicator_labels = {
            i['id']: _join_labels(structure['label'], i['label'])
            for i in structure['indicators']
        }
        for indicator in structure['indicators']:
            if indicator['column']:
                named[indicator['column']] = indicator_labels[indicator['id']]
        domain_labels = {d['name']: d['label'] for d in structure['domains']}
        for variant in (structure.get('variants') or [])[1:]:
            label = _join_labels(structure['label'], variant['label'])
            cols = variant['columns']
            named[cols.get('index')] = label
            named[cols.get('mean')] = _join_labels(
                label,
                COMPOSITE_PARTS['mean'],
            )
            named[cols.get('penalty')] = _join_labels(
                label,
                COMPOSITE_PARTS['penalty'],
            )
            for domain, column in variant['domains'].items():
                named[column] = _join_labels(
                    domain_labels.get(domain),
                    variant['label'],
                )
            for indicator, column in variant['indicators'].items():
                named[column] = _join_labels(
                    indicator_labels.get(indicator),
                    variant['label'],
                )
        for column, value in named.items():
            if column in descriptions and value:
                descriptions[column].setdefault('label', value)
    return descriptions


DASHBOARD_TYPES = ('general', 'composite', 'combined')
# the types that carry a composite index's own outputs (its raw indicator
# values, the regions' distributions, walkability's attenuation)
INDEX_TYPES = ('composite', 'combined')


def dashboard_type(config, kind=None):
    """The type of dashboard to export, and the slug it is exported under.

    ``general`` presents every configured indicator, grouped by theme;
    ``composite`` presents one composite index (``dashboard.index``, else the
    first configured) as its own dashboard, and nothing else.  The configured
    ``dashboard.type`` (default general) is exported under ``dashboard.slug``;
    the other under ``dashboard.<type>.slug``, else ``<slug>_<type>``.
    """
    configured = str(config.get('type') or 'general').lower()
    kind = str(kind or configured).lower()
    for value in (configured, kind):
        if value not in DASHBOARD_TYPES:
            raise ValueError(
                f"Unknown dashboard type '{value}' (expected one of "
                f'{DASHBOARD_TYPES}).',
            )
    if kind == configured:
        slug = config['slug']
    else:
        slug = (config.get(kind) or {}).get('slug') or (
            f'{config["slug"]}_{kind}'
        )
    return kind, slug


def select_dashboard(indicators, kind, index=None):
    """Keep the families a dashboard of this type presents.

    A general dashboard leaves out composite indices, which have a dashboard of
    their own; a composite dashboard keeps only its index, whose family also
    carries the raw values of the index's indicators; a combined one keeps
    every other family and its index.  Descriptions, themes and interventions
    are pruned to what remains.  Modifies and returns ``indicators``.
    """
    families = indicators['families']
    if kind in INDEX_TYPES:
        composites = [f for f in families if f.get('composite')]
        wanted = f'composite_{index}' if index else None
        chosen = [f for f in composites if f['id'] == wanted] or composites[:1]
        if not chosen:
            raise ValueError(
                f'A {kind} dashboard needs a composite index, and none is '
                'configured or available for this region.',
            )
    if kind == 'composite':
        kept = chosen[:1]
        indicators['themes'] = []
        indicators['interventions'] = []
    elif kind == 'combined':
        kept = [f for f in families if not f.get('composite')] + chosen[:1]
    else:
        kept = [f for f in families if not f.get('composite')]
    indicators['families'] = kept
    ids = {f['id'] for f in kept}
    for theme in indicators.get('themes') or []:
        theme['families'] = [f for f in theme['families'] if f in ids]
    indicators['themes'] = [
        t for t in indicators.get('themes') or [] if t['families']
    ]
    # an intervention is offered only for themes still presented
    themes = {t['id'] for t in indicators['themes']}
    for entry in indicators.get('interventions') or []:
        entry['themes'] = {
            k: v for k, v in (entry.get('themes') or {}).items() if k in themes
        }
    indicators['interventions'] = [
        e for e in indicators.get('interventions') or [] if e['themes']
    ]
    columns = {c for f in kept for c in f['columns']}
    indicators['descriptions'] = {
        k: v for k, v in indicators['descriptions'].items() if k in columns
    }
    used = {f['domain'] for f in kept}
    indicators['domains'] = [
        d for d in indicators.get('domains') or [] if d['id'] in used
    ]
    return indicators


TILE_MAX_COLUMNS = 60


def tile_groups(families, max_columns=TILE_MAX_COLUMNS):
    """Partition the exported columns into tile groups of bounded width.

    A feature carrying every indicator is too heavy to tile: at the zoom
    levels where a whole city fits in a few tiles, tippecanoe must drop most
    of the features to keep within its tile budget, and the map empties when
    zoomed out.  Columns are therefore tiled in groups -- one archive per scale
    and group -- by theme (families without one form an 'other' group), with a
    group wider than ``max_columns`` split into ordered chunks.  A family is
    never split across groups, so everything one view draws, and every band
    its popup reads, is in one archive.

    Returns ``{group: [columns]}`` in order, a column appearing once (in the
    group of the first family that holds it).
    """
    by_theme = {}
    for family in families:
        # names a file and a vector tile layer, so plain ASCII
        key = (
            re.sub(
                r'[^0-9a-z]+',
                '_',
                str(family.get('theme') or '').lower(),
            ).strip('_')
            or 'other'
        )
        by_theme.setdefault(key, []).append(family)
    groups = {}
    seen = set()
    for theme, members in by_theme.items():
        columns_so_far, part = [], 1
        for family in members:
            columns = [c for c in sorted(family['columns']) if c not in seen]
            seen.update(columns)
            if not columns:
                continue
            if columns_so_far and (
                len(columns_so_far) + len(columns) > max_columns
            ):
                groups[theme if part == 1 else f'{theme}_{part}'] = (
                    columns_so_far
                )
                columns_so_far, part = [], part + 1
            columns_so_far += columns
        if columns_so_far:
            groups[theme if part == 1 else f'{theme}_{part}'] = columns_so_far
    return groups


def featured_family(config, indicators):
    """The family the dashboard opens on: configured, else a composite index."""
    ids = {f['id'] for f in indicators['families']}
    featured = config.get('featured')
    if featured:
        if featured in ids:
            return featured
        print(f"  ! dashboard.featured '{featured}' is not a family; ignored")
    return next(
        (f['id'] for f in indicators['families'] if f.get('composite')),
        None,
    )


def build_indicators(r, config, available):
    """The faceted indicator vocabulary the dashboard navigates."""
    # Applied here rather than family by family: everything below is built from
    # `available`, so a hidden column simply never enters the vocabulary and
    # cannot reappear as a variable, a band or a dictionary entry.  Used for
    # outputs that are real but not worth navigating -- the GUHVI sub-indicators
    # include a subnational HDI and an infant mortality rate that are constant
    # across a single city and map as one flat colour.
    hidden = set(config.get('hide') or [])
    if hidden:
        available = {c for c in available if c not in hidden}
    ped, cyc = _mode_configs(r)
    families = {}
    for family in destination_families(r, ped, cyc, available):
        families[family['id']] = family

    for resolved, networks in (
        (ped, [WALK]),
        (cyc, (cyc or {}).get('networks', [])),
    ):
        if resolved is None or not networks:
            continue
        ac_names, ac_labels = activity_centre_names(resolved['config'])
        _merge_families(
            families,
            _derived_families(
                resolved,
                networks,
                available,
                ac_names,
                'activity_centre',
                'Activity centres',
                overlay_for=lambda n: {'layer': n},
                labels=ac_labels,
            ),
        )
        all_names, all_labels = combined_access_names(
            resolved['config'],
            resolved['specs'],
        )
        _merge_families(
            families,
            _derived_families(
                resolved,
                networks,
                available,
                all_names,
                'combined',
                'Combined access',
                labels=all_labels,
            ),
        )

    ordered = list(families.values())
    ordered += diversity_families(r, ped, available)
    ordered += catchment_families(available)
    core = core_access_family(available)
    if core:
        ordered.append(core)
    ordered += standalone_families(available)
    ordered += walkability_variant_families(r, available)
    ordered += linkage_families(r, available)
    ordered += composite_families(r, available, config)

    themes = resolve_themes(config)
    assign_themes(themes, ordered)
    # provenance, from the region's own configuration
    notes = config.get('notes') or {}
    ped, _cyc = _mode_configs(r)
    specs = {s['name']: s for s in (ped or {}).get('specs', [])}
    for family in ordered:
        spec = specs.get(family['id'])
        if spec:
            sources = destination_provenance(r, spec)
            if sources:
                family['sources'] = sources
        note = notes.get(family['id']) or (
            notes.get((family.get('overlay') or {}).get('layer'))
        )
        if note:
            family['note'] = _labels(note, '')
    crosswalk = load_crosswalk(config.get('crosswalk'))
    descriptions = {}
    for family in ordered:
        family.setdefault('group', 'destination')
        family.setdefault('direction', HIGHER)
        columns = _family_columns(family)
        uli = {}
        for column in sorted(columns):
            category, text = describe_variable(column)
            # units come from describe_units() rather than the region's
            # generated dictionary CSV: that file was written whenever
            # generate() last ran, so a correction to the units would not reach
            # the dashboard until the whole region was regenerated
            units, statistic = describe_units(column)
            descriptions[column] = {
                'category': category,
                'en': text,
                'units': units,
                'statistic': statistic,
            }
            entry = crosswalk.get(column)
            if entry:
                uli.setdefault(entry['measure'], entry)
                if entry.get('direction'):
                    descriptions[column]['direction'] = entry['direction']
        if uli:
            family['uli'] = list(uli.values())
        family['columns'] = sorted(columns)

    apply_linked_metadata(ordered, descriptions)
    for family in ordered:
        labels = family.pop('variable_labels', None) or {}
        units = family.pop('variable_units', None) or {}
        for column, label in labels.items():
            if column in descriptions:
                descriptions[column].setdefault('label', label)
                if column in units:
                    descriptions[column]['units'] = units[column]
    apply_labels(config, ordered, descriptions)
    apply_composite_labels(ordered, descriptions)
    compose_spanish_descriptions(ordered, descriptions)
    interventions = load_interventions(config.get('interventions'), themes)

    domains = []
    for family in ordered:
        if family['domain'] not in domains:
            domains.append(family['domain'])
    # only the networks something was actually measured over: a region that ran
    # two of the four cycling measures should not offer the other two at all
    used = {
        network
        for family in ordered
        for measure in family['measures'].values()
        for network in (measure.get('networks') or {})
    }
    return {
        # the sentences the viewer composes, as templates.  Emitted rather than
        # written into the JavaScript so that the wording can be revised in the
        # site's own text.json without a code change or a re-export; see
        # DEFAULT_TEXT for the placeholders each one may use.
        'text': DEFAULT_TEXT,
        # the workshop's own typology, and the primary grouping in the viewer;
        # the analysis-derived domains remain as a secondary one
        'themes': [
            {k: v for k, v in theme.items() if k != 'keys'}
            for theme in themes.values()
        ],
        'interventions': interventions,
        'rules': analysis_rules(r),
        'domains': [
            {'id': d, 'label': {'es': ES.get(d, d), 'en': d}} for d in domains
        ],
        'networks': {k: v for k, v in NETWORKS.items() if k in used},
        'measures': {k: v for k, v in MEASURE_META.items()},
        'families': ordered,
        'descriptions': descriptions,
    }


# ---------------------------------------------------------------------------
# Layer export

# Rounding, chosen per column so that tiles dedup well without losing meaning:
# percentages and metres are integers, indices on 0-1 keep three decimals, and
# anything else keeps two.
INTEGER_PREFIXES = (
    'pct_',
    'avg_walk_dist_',
    'avg_cycle_dist_',
    'avg_cycle_extra_',
    'intersection_count',
)
# pop_est is deliberately not an integer: where a scale's population is an
# assumed figure per unit rather than a count -- 3.2 residents per Condesa lot
# -- rounding is the difference between reporting the assumption and losing it,
# and it made the lot popup contradict the note printed beside it.  Whole
# numbers are unaffected: a grid cell of 47 people rounds to 47.0 and tiles
# dedup on it just the same.
INDEX_PREFIXES = (
    'avg_diversity_walk_',
    'avg_richness_walk_',
    'urban_heat_',
    # heat attenuated walkability is a rank on 0-1, and NDVI a ratio on -1-1
    'walk_idx_',
    'ext_ndvi',
)


def round_expression(column, source='i', index_decimals=None):
    """SQL that rounds a column to a sensible precision, aliased to itself.

    ``index_decimals`` rounds composite index scores (``index_*``) to the
    precision they are displayed at: a composite dashboard's features carry
    little else, and a score to one decimal dedups far better in the tiles.
    """
    qualified = f'{source}."{column}"'
    if index_decimals is not None and column.startswith('index_'):
        return (
            f'ROUND({qualified}::numeric, {int(index_decimals)})::float8 '
            f'AS "{column}"'
        )
    if column.startswith(INTEGER_PREFIXES):
        return f'ROUND({qualified}::numeric)::int AS "{column}"'
    if column.startswith(INDEX_PREFIXES):
        return f'ROUND({qualified}::numeric, 3)::float8 AS "{column}"'
    return f'ROUND({qualified}::numeric, 2)::float8 AS "{column}"'


def write_layer(gdf, path):
    """Write a layer as newline-delimited GeoJSON in EPSG:4326."""
    gdf = gdf.to_crs(4326)
    try:
        gdf.to_file(path, driver='GeoJSONSeq', COORDINATE_PRECISION=6)
    except Exception:
        gdf.to_file(path, driver='GeoJSONSeq')
    return len(gdf)


def grid_raster(gdf, tiles, outdir, layer):
    """Write a regular grid's values as a raster the viewer can smooth.

    The 100 m population grid is a regular raster in the project's
    projection, polygonised: each cell an axis-aligned square of one size.  So
    it is written back as one: each area's cell (row-major, from the upper
    left) as ``<layer>.cells.i32``, and each tile group's columns as
    ``<layer>__<group>.f32``, column by column in the area order (NaN for no
    value), little-endian.  The four corners of the grid's extent are given in
    longitude and latitude, for placing the smoothed image on the map.

    Returns the scale's ``raster`` entry, or None where the areas are not a
    regular grid.
    """
    bounds = gdf.geometry.bounds
    widths = bounds['maxx'] - bounds['minx']
    heights = bounds['maxy'] - bounds['miny']
    cell = float(widths.median())
    if (
        not cell > 0
        or (widths - cell).abs().max() > cell * 1e-3
        or (heights - cell).abs().max() > cell * 1e-3
    ):
        return None
    x0 = float(bounds['minx'].min())
    y0 = float(bounds['maxy'].max())
    cols = ((bounds['minx'] - x0) / cell).round()
    rows = ((y0 - bounds['maxy']) / cell).round()
    if ((bounds['minx'] - x0) / cell - cols).abs().max() > 1e-3:
        return None
    nx = int(cols.max()) + 1
    ny = int(rows.max()) + 1
    index = (rows * nx + cols).astype('int32').to_numpy()
    index.astype('<i4').tofile(f'{outdir}/{layer}.cells.i32')
    groups = {}
    for group, tile in tiles.items():
        values = np.stack(
            [
                pd.to_numeric(gdf[c], errors='coerce').to_numpy(
                    dtype='float32',
                )
                for c in tile['columns']
            ],
        )
        name = f'{layer}__{group}.f32'
        values.astype('<f4').tofile(f'{outdir}/{name}')
        groups[group] = {'file': name, 'columns': tile['columns']}
    import geopandas as gpd
    from shapely.geometry import Point

    corners = gpd.GeoSeries(
        [
            Point(x0, y0),
            Point(x0 + nx * cell, y0),
            Point(x0 + nx * cell, y0 - ny * cell),
            Point(x0, y0 - ny * cell),
        ],
        crs=gdf.crs,
    ).to_crs(4326)
    return {
        'cell': cell,
        'nx': nx,
        'ny': ny,
        'areas': int(len(gdf)),
        # upper left, upper right, lower right, lower left, as MapLibre's
        # image sources take them
        'corners': [[round(p.x, 6), round(p.y, 6)] for p in corners],
        'index': f'{layer}.cells.i32',
        'groups': groups,
    }


CONTEXT_COLUMNS = [
    'pop_est',
    'area_sqkm',
    'pop_per_sqkm',
    'intersection_count',
    'intersections_per_sqkm',
]


# an area is exported where a point on its surface lies in the study region
WITHIN_REGION = (
    'ST_Intersects(ST_PointOnSurface(b.geom), '
    '(SELECT ST_Union(geom) FROM urban_study_region))'
)


def scale_query(scale, resolved, wanted, with_geom=True, index_decimals=None):
    """SQL selecting one scale's areas within the study region, every one kept.

    Areas that received no source units are deleted from the indicator table by
    _12_aggregation, so the boundaries are the left side of the join: an area
    with no result is exported with null values and drawn as 'no data' rather
    than silently missing from the map.  Areas outside the study region are
    not exported at all -- a boundary layer such as the census manzanas covers
    the whole municipality, and its areas beyond the region were never sampled,
    so on the map they read as unmeasured parts of the city.  An area belongs
    to the region where a point on its surface lies within it.
    """
    id_column = scale['id']
    joined = scale['boundaries'] != scale['table']
    source = 'i' if joined else 'b'
    select = [f'b."{id_column}"::text AS area_id']
    for column in scale.get('keep_columns') or []:
        select.append(f'b."{column}"')
    for canonical in wanted:
        physical, _ = resolved[canonical]
        expression = round_expression(physical, source, index_decimals)
        if physical != canonical:
            expression = f'{expression.rsplit(" AS ", 1)[0]} AS "{canonical}"'
        select.append(expression)
    if with_geom:
        select.append('b.geom')
    # ordered so that repeated exports are byte-reproducible: without it the
    # row order is whatever the planner returns, which changes both the feature
    # order in the layer and the weighted quantiles computed from it
    order_by = f'ORDER BY b."{id_column}"'
    where = f'WHERE {WITHIN_REGION}'
    if not joined:
        return (
            f'SELECT {", ".join(select)} FROM {scale["table"]} b '
            f'{where} {order_by}'
        )
    return (
        f'SELECT {", ".join(select)} FROM {scale["boundaries"]} b '
        f'LEFT JOIN {scale["table"]} i '
        f'ON i."{id_column}"::text = b."{id_column}"::text '
        f'{where} {order_by}'
    )


def export_scale(
    r,
    scale,
    vocabulary_columns,
    outdir,
    write=True,
    groups=None,
    index_decimals=None,
    raster=False,
):
    """Write one scale's layers; returns its manifest entry.

    With ``groups`` (see :func:`tile_groups`) the scale is written as one layer
    per group, ``scale_<key>__<group>``, each holding the area id, the kept
    boundary columns, the context columns and that group's indicators.
    """
    resolved = canonical_columns(r, scale['table'])
    # de-duplicated: the context variables are also offered as a family, so
    # they appear in the vocabulary as well as in CONTEXT_COLUMNS
    wanted = [
        c
        for c in dict.fromkeys(list(vocabulary_columns) + CONTEXT_COLUMNS)
        if c in resolved
    ]
    # keep_columns are read from the boundaries, so only those actually there
    boundary_columns = set(table_columns(r, scale['boundaries']))
    scale = dict(
        scale,
        keep_columns=[
            c
            for c in (scale.get('keep_columns') or [])
            if c in boundary_columns
        ],
    )
    if scale['id'] not in boundary_columns:
        print(
            f'  ! {scale["key"]}: id column "{scale["id"]}" not on '
            f'{scale["boundaries"]}; falling back to ogc_fid',
        )
        scale['id'] = 'ogc_fid' if 'ogc_fid' in boundary_columns else 'fid'
    layer = f'scale_{scale["key"]}'
    context = [c for c in CONTEXT_COLUMNS if c in wanted]
    fixed = ['area_id', *scale['keep_columns'], *context]
    tiles = {}
    for group, columns in (groups or {}).items():
        present = [c for c in columns if c in wanted and c not in context]
        if present:
            tiles[group] = {
                'layer': f'{layer}__{group}',
                'file': f'{layer}__{group}.geojsonl',
                'columns': present,
            }
    if not write:
        # the geometry is unchanged, so only the feature count is needed and it
        # can be counted in the database instead of serialised out again
        features = int(
            r.get_df(
                f'SELECT count(*) AS n FROM {scale["boundaries"]} b '
                f'WHERE {WITHIN_REGION}',
            )['n'].iloc[0],
        )
    else:
        gdf = r.get_gdf(
            scale_query(
                scale,
                resolved,
                wanted,
                index_decimals=index_decimals,
            ),
        )
        if gdf is None:
            print(f'  ! {scale["key"]}: query failed, skipping')
            return None, None
        if tiles:
            for tile in tiles.values():
                features = write_layer(
                    gdf[[*fixed, *tile['columns'], gdf.geometry.name]],
                    f'{outdir}/{tile["file"]}',
                )
        else:
            features = write_layer(gdf, f'{outdir}/{layer}.geojsonl')
        if raster and tiles:
            if gdf.crs is None:
                gdf = gdf.set_crs(r.config['crs']['srid'])
            raster = grid_raster(gdf, tiles, outdir, layer)
            if raster is None:
                print(f'  ! {scale["key"]}: not a regular grid; no raster')
    weighted = sorted(c for c in wanted if resolved[c][1] == 'weighted')
    print(
        f'  {layer}: {features} areas, {len(wanted)} indicators'
        + (
            f' in {len(tiles)} tile groups (widest '
            f'{max(len(t["columns"]) for t in tiles.values())})'
            if tiles
            else ''
        ),
        flush=True,
    )
    return scale, {
        'key': scale['key'],
        'layer': layer,
        'file': f'{layer}.geojsonl',
        'features': features,
        'label': scale['label'],
        'id': scale['id'],
        'weight': scale.get('weight'),
        'source': scale.get('source'),
        'keep_columns': scale['keep_columns'],
        'columns': wanted,
        # the layers the scale is tiled as, by group: each archive holds the
        # fixed columns and its group's indicators
        'tiles': tiles,
        # a regular grid's values as a raster, for a smoothed surface
        **({'raster': raster} if isinstance(raster, dict) else {}),
        # columns whose value came from a population-weighted aggregation
        # column rather than one carrying the canonical name: the same
        # quantity, differently derived, and the info panel says so
        'weighted_columns': weighted,
    }


# Edge attributes worth carrying into the network overlay's popups.
EDGE_COLUMNS = [
    'osmid',
    'name',
    'highway',
    'maxspeed_kmh',
    'adt',
    'bike_facility',
    'lvl_traf_stress',
    'bike_permitted',
    'foot_dismount',
]


def _previous_manifest(outdir):
    """The manifest the last export wrote here, or {}."""
    path = f'{outdir}/manifest.json'
    if not os.path.exists(path):
        return {}
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def _previous_layers(outdir):
    """The layer index recorded by the last full export, for --no-layers."""
    path = f'{outdir}/manifest.json'
    if not os.path.exists(path):
        print(
            '  ! --no-layers with no previous manifest; layers will be empty',
        )
        return {}
    with open(path, encoding='utf-8') as f:
        return json.load(f).get('layers') or {}


def export_context_layers(r, indicators, outdir):
    """Write the overlay layers that explain a result.

    The points an indicator was measured to, the network it was measured over,
    and the region boundaries.
    """
    tables = set(r.get_tables())
    layers = {}

    def record(name, gdf, **extra):
        if gdf is None or len(gdf) == 0:
            return
        features = write_layer(gdf, f'{outdir}/{name}.geojsonl')
        # geometry type is recorded rather than inferred from the layer name:
        # the viewer needs to know whether to draw circles or lines, and open
        # space "entry point" layers are lines despite being conceptually points
        kinds = set(gdf.geom_type.dropna().unique())
        geometry = (
            'line'
            if any('Line' in k for k in kinds)
            else 'polygon' if any('Polygon' in k for k in kinds) else 'point'
        )
        layers[name] = dict(
            {
                'file': f'{name}.geojsonl',
                'features': features,
                'geometry': geometry,
            },
            **extra,
        )
        print(f'  {name}: {features} {geometry} features', flush=True)

    # every point layer any indicator refers to
    referenced = {}
    for family in indicators['families']:
        overlay = family.get('overlay')
        if overlay and overlay.get('layer'):
            referenced.setdefault(overlay['layer'], set()).update(
                overlay.get('dest_names') or [],
            )

    if 'destinations' in referenced and 'destinations' in tables:
        names = sorted(referenced.pop('destinations'))
        record(
            'destinations',
            r.get_gdf(
                'SELECT dest_name, dest_name_full, geom FROM destinations'
                f" WHERE dest_name IN ({', '.join(repr(n) for n in names)})",
            ),
            dest_names=names,
        )
    for layer in sorted(referenced):
        if layer not in tables:
            continue
        record(layer, r.get_gdf(f'SELECT geom FROM {layer}'))

    columns = [c for c in EDGE_COLUMNS if c in set(table_columns(r, 'edges'))]
    record(
        'network',
        r.get_gdf(
            f'SELECT {", ".join(columns)}, '
            'ROUND(length::numeric)::int AS length_m, geom FROM edges',
        ),
    )
    record(
        'boundary',
        r.get_gdf(
            'SELECT study_region, area_sqkm, pop_est, geom '
            'FROM urban_study_region',
        ),
    )
    buffer_table = r.config.get('buffered_urban_study_region')
    if buffer_table in tables:
        record('buffer', r.get_gdf(f'SELECT geom FROM {buffer_table}'))
    return layers


# ---------------------------------------------------------------------------
# Distributions
#
# The viewer never reads tile attributes to decide a colour scale or draw a
# histogram: both come from these summaries, so a shared classification across
# two panes showing different scales costs one small fetch instead of a scan of
# every rendered feature.

HISTOGRAM_BINS = 12
# a cell/area counts as having access at a band when at least half its sample
# points do -- the same rule the choropleth colours by
BAND_THRESHOLD = 50


def weighted_quantile(values, weights, q):
    # lexsort, not argsort: rows arrive in whatever order the database returned
    # them, and an unstable sort permutes tied values among themselves.  The
    # tied values are identical, but their weights are not, so the cumulative
    # weight at the start of a run of ties -- and hence the interpolation from
    # the value before it -- would otherwise vary between runs.  Sorting on
    # (value, weight) fixes the result to the rows themselves.
    order = np.lexsort((weights, values))
    v, w = values[order], weights[order]
    cw = np.cumsum(w)
    if cw[-1] <= 0:
        return None
    return float(np.interp(q * cw[-1], cw, v))


def _number(value):
    if value is None or pd.isna(value):
        return None
    value = float(value)
    if value != value or value in (float('inf'), float('-inf')):
        return None
    return round(value, 4)


def column_ranges(r, exported):
    """The min and max of every exported column, across every scale.

    Histograms are binned on these shared edges rather than on each scale's own
    range, so that two panes showing different scales can be drawn as grouped
    bars: bars that do not share their bins cannot be compared side by side.
    """
    ranges = {}
    for entry, scale in exported:
        columns = entry['columns']
        if not columns:
            continue
        resolved = canonical_columns(r, scale['table'])
        select = ', '.join(
            f'min("{resolved[c][0]}") AS "min_{i}", '
            f'max("{resolved[c][0]}") AS "max_{i}"'
            for i, c in enumerate(columns)
        )
        row = r.get_df(f'SELECT {select} FROM {scale["table"]}').iloc[0]
        for i, column in enumerate(columns):
            low, high = row[f'min_{i}'], row[f'max_{i}']
            if low is None or pd.isna(low):
                continue
            previous = ranges.get(column)
            low, high = float(low), float(high)
            ranges[column] = (
                (low, high)
                if previous is None
                else (min(previous[0], low), max(previous[1], high))
            )
    return ranges


COMPOSITE_SPREAD = (0.05, 0.95)


def composite_spreads(r, exported, indicators, quantiles=COMPOSITE_SPREAD):
    """The middle 90% of each composite index and domain score.

    The shared classes of a composite index are sized from this rather than
    from its range (see ``_composite_index.composite_classes``).  Each column's
    span is taken from the scale with the most values for it -- the finest --
    rather than widened over every scale as :func:`column_ranges` is: a small
    exported area of unusual places (a fringe development, say) would
    otherwise set the classes for the whole city.
    """
    headline = set()
    for family in indicators['families']:
        structure = family.get('composite')
        if structure:
            headline.add(structure['columns']['index'])
            headline.update(
                d['column'] for d in structure['domains'] if d.get('column')
            )
    low_q, high_q = quantiles
    spreads, counts = {}, {}
    for entry, scale in exported:
        columns = [c for c in entry['columns'] if c in headline]
        if not columns:
            continue
        resolved = canonical_columns(r, scale['table'])
        columns = [c for c in columns if c in resolved]
        if not columns:
            continue
        select = ', '.join(
            f'percentile_cont({low_q}) WITHIN GROUP '
            f'(ORDER BY "{resolved[c][0]}") AS "low_{i}", '
            f'percentile_cont({high_q}) WITHIN GROUP '
            f'(ORDER BY "{resolved[c][0]}") AS "high_{i}", '
            f'count("{resolved[c][0]}") AS "n_{i}"'
            for i, c in enumerate(columns)
        )
        row = r.get_df(f'SELECT {select} FROM {scale["table"]}').iloc[0]
        for i, column in enumerate(columns):
            low, high = row[f'low_{i}'], row[f'high_{i}']
            if low is None or pd.isna(low):
                continue
            n = int(row[f'n_{i}'])
            if column not in counts or n > counts[column]:
                counts[column] = n
                spreads[column] = (float(low), float(high))
    return spreads


# Figures a conceptual model may be given as, by extension, and how each is
# shown: an image, or a document in a frame.
CONCEPTUAL_MODEL_TYPES = {
    '.svg': 'image',
    '.png': 'image',
    '.jpg': 'image',
    '.jpeg': 'image',
    '.webp': 'image',
    '.pdf': 'document',
}
# report language names to the site's language codes, where the reporting
# configuration's language sheet cannot be read
LANGUAGE_CODES = {'english': 'en', 'spanish': 'es', 'español': 'es'}


def language_codes(r):
    """Report language names to language codes (``English`` -> ``en``).

    Read from the ``language_code`` row of the reporting configuration's
    languages sheet, which is what the PDF reports use.
    """
    path = (r.config.get('reporting') or {}).get('configuration')
    try:
        sheet = pd.read_excel(path, sheet_name='languages').set_index('name')
        row = sheet.loc['language_code']
        return {
            str(k): str(v).strip()
            for k, v in row.items()
            if k != 'role' and isinstance(v, str) and v.strip()
        }
    except Exception:
        return {}


def language_code(name, codes=None):
    """A language's code, from the sheet or, failing that, its name."""
    if name in (codes or {}):
        return codes[name]
    first = str(name).split(' - ')[0].strip().lower()
    return LANGUAGE_CODES.get(first, first[:2])


def conceptual_models(r, outdir):
    """Copy each language's conceptual model figure into the export.

    Configured per report language as ``reporting.languages.<Language>.
    conceptual_model``: a path relative to process/data (or absolute), or
    ``{file, caption, alt}``.  Each is written as ``conceptual_model_<code>
    .<ext>``.  Returns ``{code: {file, type, caption, alt}}`` for the manifest;
    a missing file or unsupported type is reported and left out.
    """
    languages = (r.config.get('reporting') or {}).get('languages') or {}
    codes = None
    models = {}
    for name, settings in languages.items():
        entry = (settings or {}).get('conceptual_model')
        if not entry:
            continue
        if isinstance(entry, str):
            entry = {'file': entry}
        source = str(entry['file'])
        if not os.path.isabs(source):
            source = _data_path(source)
        extension = os.path.splitext(source)[1].lower()
        if extension not in CONCEPTUAL_MODEL_TYPES:
            supported = ', '.join(sorted(CONCEPTUAL_MODEL_TYPES))
            print(
                f'  ! conceptual model for {name}: unsupported type '
                f"'{extension}' (expected one of {supported})",
            )
            continue
        if not os.path.exists(source):
            print(f'  ! conceptual model for {name}: {source} not found')
            continue
        if codes is None:
            codes = language_codes(r)
        code = language_code(name, codes)
        target = f'conceptual_model_{code}{extension}'
        shutil.copyfile(source, os.path.join(outdir, target))
        models[code] = {
            'file': target,
            'type': CONCEPTUAL_MODEL_TYPES[extension],
            'caption': entry.get('caption'),
            'alt': entry.get('alt'),
        }
    if models:
        print(f'  Conceptual models: {", ".join(sorted(models))}')
    return models


def histogram_edges(span):
    """Bin edges for a column, from its range across every scale."""
    low, high = span
    if high > low:
        return np.linspace(low, high, HISTOGRAM_BINS + 1)
    return np.array([low, low + 1.0])


# ---------------------------------------------------------------------------
# Classification
#
# Where the map's colour classes come from.  This used to live in the viewer,
# which meant the histogram could only be binned on its own uniform edges and
# the swatches beneath it never lined up with the bars.  Here the classes and
# the population share falling in each are computed together, from the same
# data, so the legend *is* the chart's axis.
#
# A column's classes come from, in order of precedence:
#
#   1. the region's ``dashboard.breaks`` entry for that column;
#   2. its declared units, where the data actually occupies them;
#   3. its range, cut into round numbers.

# How much of a declared range the data must occupy for that range to be worth
# classifying over.  NDVI is declared "-1 to 1" and holds 0.03 to 0.35:
# legitimate, but classifying over the declared range puts every area in one
# class and draws the same flat map as an outright wrong unit would.
MIN_OCCUPANCY = 0.25
DEFAULT_LOG2_CLASSES = 7
# Default classes for a whole measure, where its range says less about it than
# its meaning does.  A mean count of reachable establishments runs from 0 to 64
# for produce within a kilometre, so classifying from the range puts nearly
# every area in the bottom class; what a reader wants to know is whether there
# are none, one, a couple, or plenty.  Keyed by measure rather than by units
# because `intersection_count` is also a count and is nothing like this.
MEASURE_BREAKS = {
    # closed at the bottom -- a count of reachable places cannot be negative,
    # and an empty "< 0" class would take a column of the legend to say so
    'count': {
        'edges': [0, 1, 2, 3, 4, 5],
        'open_low': False,
        'open_high': True,
    },
}
# A units string may state its own range -- "index 0-1", "index -1 to 1",
# "class 1-5".  Anything else ("index (sum of z-scores)") states nothing.
_DECLARED_RANGE = re.compile(
    r'(-?\d+(?:\.\d+)?)\s*(?:-|to|–)\s*(-?\d+(?:\.\d+)?)',
)
# round distances to step up through until the data fits
_METRE_LADDERS = [
    [0, 100, 250, 500, 750, 1000],
    [0, 250, 500, 1000, 1500, 2000],
    [0, 500, 1000, 2000, 3000, 5000],
    [0, 1000, 2000, 5000, 10000, 20000],
]


def declared_range(units):
    """The range a units string states, or None if it states none."""
    text = str(units or '').lower()
    if 'percent' in text:
        return (0.0, 100.0)
    match = _DECLARED_RANGE.search(text)
    if not match:
        return None
    low, high = float(match.group(1)), float(match.group(2))
    return (low, high) if high > low else None


def _nice_step(rough):
    """A round step at roughly the requested size: 1, 2, 2.5 or 5 x 10^n."""
    if not rough > 0:
        return 1.0
    magnitude = 10 ** math.floor(math.log10(rough))
    normalised = rough / magnitude
    for limit, step in ((1, 1), (2, 2), (2.5, 2.5), (5, 5)):
        if normalised <= limit:
            return step * magnitude
    return 10 * magnitude


def nice_breaks(span, n=5):
    """Round edges spanning a range, so the legend reads in whole numbers."""
    low, high = span
    if not high > low:
        return [low, low + 1.0]
    step = _nice_step((high - low) / n)
    start = math.floor(low / step) * step
    edges = []
    while start + step * len(edges) <= high + step * 1e-9:
        edges.append(round(start + step * len(edges), 10))
    edges.append(round(start + step * len(edges), 10))
    return edges


def _equal_interval(span, n=5):
    low, high = span
    if not high > low:
        return [low, low + 1.0]
    step = (high - low) / n
    return [round(low + step * i, 10) for i in range(n + 1)]


def resolve_targets(r, config):
    """Targets a column can be measured against, keyed by column.

    Seeded from the GHSCI reference configuration's own ``thresholds`` block
    (``indicators.yml``), which already records exactly this -- a field, a
    criteria and a relationship -- so the neighbourhood density targets are
    present in every region without configuring anything.  The region's
    ``dashboard.thresholds`` block overrides and extends it, which is how a
    target agreed after the workshop gets added without touching code.
    """
    targets = {}
    try:
        declared = (r.indicators or {}).get('report', {}).get('thresholds', {})
    except AttributeError:
        declared = {}
    for title, spec in (declared or {}).items():
        field = (spec or {}).get('field')
        criteria = (spec or {}).get('criteria')
        if not field or criteria is None:
            continue
        targets[field] = {
            'criteria': float(criteria),
            'relationship': str(spec.get('relationship') or '>='),
            'title': _labels(spec.get('title') or title, field),
        }
    for column, spec in (config.get('thresholds') or {}).items():
        spec = spec if isinstance(spec, dict) else {'criteria': spec}
        criteria = spec.get('criteria')
        if criteria is None:
            targets.pop(column, None)
            continue
        entry = dict(targets.get(column) or {})
        entry['criteria'] = float(criteria)
        entry['relationship'] = str(
            spec.get('relationship') or entry.get('relationship') or '>=',
        )
        if spec.get('title') or not entry.get('title'):
            entry['title'] = _labels(spec.get('title'), column)
        targets[column] = entry
    return targets


def _log2_edges(anchor, span, classes):
    """Edges at ``anchor x 2^k``, covering the data, anchored on the target.

    A doubling ladder is how an unbounded, heavily skewed density reads: drawn
    as equal-width classes it *is* a log axis, without transforming any value.
    The ladder is hung off the target so that the target is always an edge and
    the top class means exactly "meets it".  The lowest class stays open, which
    is what lets a column holding genuine zeros -- both neighbourhood densities
    do -- be shown on a log scale at all.
    """
    if not anchor > 0:
        return None
    low, high = span
    steps = max(2, int(classes) - 1)
    # the highest edge still below the observed maximum, so the top class is
    # populated rather than an empty "beyond everything"
    top = 0
    while anchor * (2 ** (top + 1)) < high:
        top += 1
        if top > 8:
            break
    edges = [anchor * (2 ** (top - i)) for i in range(steps)][::-1]
    # trim edges the data never reaches down to, but keep one below the minimum
    while len(edges) > 2 and edges[1] <= low:
        edges.pop(0)
    # halving a target like 5700 lands on 178.125, and a legend does not want
    # three decimal places of people per square kilometre
    return [round(e) if abs(e) >= 100 else round(e, 2) for e in edges]


def _configured_breaks(spec, span, target):
    """Classes from a ``dashboard.breaks`` entry, or None if it defines none."""
    if isinstance(spec, (list, tuple)):
        # a bare list is interior breaks: the tails are grouped
        return {
            'kind': 'classes',
            'edges': [float(v) for v in spec],
            'open_low': True,
            'open_high': True,
        }
    if not isinstance(spec, dict):
        return None
    if spec.get('categories'):
        return {
            'kind': 'categories',
            'values': [float(v) for v in spec['categories']],
        }
    if str(spec.get('scale') or '').lower() in ('log2', 'log'):
        anchor = spec.get('anchor')
        if anchor is None and target:
            anchor = target.get('criteria')
        edges = _log2_edges(
            float(anchor or 0),
            span,
            spec.get('classes') or DEFAULT_LOG2_CLASSES,
        )
        if not edges:
            return None
        return {
            'kind': 'classes',
            'edges': edges,
            'open_low': True,
            'open_high': True,
            'scale': 'log2',
        }
    if spec.get('edges'):
        closed = bool(spec.get('closed'))
        return {
            'kind': 'classes',
            'edges': [float(v) for v in spec['edges']],
            'open_low': not closed and bool(spec.get('open_low', True)),
            'open_high': not closed and bool(spec.get('open_high', True)),
        }
    return None


def class_breaks(units, span, spec=None, target=None):
    """The colour classes for one column.

    ``span`` is its (min, max) across every exported scale, so one colour means
    one thing on both panes however they differ in scale.
    """
    configured = _configured_breaks(spec, span, target)
    if configured:
        return configured
    text = str(units or '').lower()
    if 'metre' in text:
        high = span[1]
        ladder = next(
            (rungs for rungs in _METRE_LADDERS if high <= rungs[-1]),
            _METRE_LADDERS[-1],
        )
        # closed: the ladder was chosen so the data fits under its top edge,
        # and an open class above it would always be empty
        return {
            'kind': 'classes',
            'edges': [float(v) for v in ladder],
            'open_low': False,
            'open_high': False,
        }
    declared = declared_range(text)
    if declared:
        fits = span[0] >= declared[0] - 1e-9 and span[1] <= declared[1] + 1e-9
        if fits and 'percent' in text:
            # a percentage is always classified 0-100, however narrow the data:
            # it is the one scale every reader already knows, and holding it
            # fixed keeps colours comparable from one indicator to the next
            return {
                'kind': 'classes',
                'edges': [0.0, 20.0, 40.0, 60.0, 80.0, 100.0],
                'open_low': False,
                'open_high': False,
            }
        if fits:
            occupancy = (span[1] - span[0]) / (declared[1] - declared[0])
            if occupancy >= MIN_OCCUPANCY:
                return {
                    'kind': 'classes',
                    'edges': _equal_interval(declared, 5),
                    'open_low': False,
                    'open_high': False,
                }
        elif not fits:
            # the data escapes its declared units: say so, then classify honestly
            print(
                f'    values {span[0]}-{span[1]} fall outside the declared '
                f'units "{units}"; classifying from the data instead',
            )
    # closed for the same reason: nice_breaks() spans the data by construction
    return {
        'kind': 'classes',
        'edges': nice_breaks(span, 5),
        'open_low': False,
        'open_high': False,
    }


def class_count(breaks):
    """How many classes a break definition describes."""
    if not breaks:
        return 0
    if breaks['kind'] == 'categories':
        return len(breaks['values'])
    return (
        len(breaks['edges'])
        - 1
        + int(breaks.get('open_low', False))
        + int(breaks.get('open_high', False))
    )


def class_index(values, breaks):
    """The class each value falls in, as an integer array."""
    if breaks['kind'] == 'categories':
        wanted = np.asarray(breaks['values'], dtype=float)
        idx = np.full(len(values), -1, dtype=int)
        for i, value in enumerate(wanted):
            idx[np.isclose(values, value)] = i
        return idx
    edges = np.asarray(breaks['edges'], dtype=float)
    raw = np.digitize(values, edges, right=False)
    low = bool(breaks.get('open_low', False))
    high = bool(breaks.get('open_high', False))
    if low and high:
        return raw
    if low:
        return np.clip(raw, 0, len(edges) - 1)
    if high:
        return np.clip(raw - 1, 0, len(edges) - 1)
    return np.clip(raw - 1, 0, len(edges) - 2)


def measure_of_column(indicators):
    """``{column: measure key}``, for measure-wide classification defaults."""
    by_column = {}
    for family in indicators['families']:
        for key, measure in family['measures'].items():
            for column in _measure_columns(measure):
                by_column.setdefault(column, key)
    return by_column


def _measure_columns(measure):
    """Every column one measure entry names, however it is nested."""
    found = []

    def walk(node):
        if isinstance(node, str):
            found.append(node)
        elif isinstance(node, dict):
            for value in node.values():
                walk(value)

    for key in ('networks', 'groups', 'variables'):
        walk(measure.get(key))
    return found


def composite_class_breaks(indicators, ranges, spreads=None):
    """The shared diverging classes of every composite index's scores."""
    from _composite_index import composite_classes

    out = {}
    for family in indicators['families']:
        structure = family.get('composite')
        if not structure:
            continue
        classes = composite_classes(structure, ranges, spreads=spreads)
        out.update(classes)
        # every variant shares the base index's classes, so that switching
        # between them changes colours only where scores change
        shared = classes.get(structure['columns']['index'])
        if not shared:
            continue
        for variant in structure.get('variants') or []:
            for group in ('columns', 'domains', 'indicators'):
                for key, column in (variant.get(group) or {}).items():
                    if key == 'penalty' or column not in ranges:
                        continue
                    out.setdefault(column, dict(shared))
    return out


def all_class_breaks(indicators, ranges, targets, configured, spreads=None):
    """Break definitions for every column that has a range, keyed by column.

    Precedence: the region's own ``dashboard.breaks`` entry, then the shared
    classes of a composite index's scores, then a censored distance's own
    bands (so no class lies beyond what was searched), then a default for the
    whole measure, then the column's declared units, then its range.  A
    configured entry for a composite score keeps the index's diverging ramp.
    """
    configured = configured or {}
    descriptions = indicators['descriptions']
    by_measure = measure_of_column(indicators)
    composite = composite_class_breaks(indicators, ranges, spreads)
    censored = indicators.get('censored') or {}
    breaks = {}
    for column, span in ranges.items():
        described = descriptions.get(column) or {}
        spec = configured.get(column)
        if spec is None and column in composite:
            breaks[column] = composite[column]
            continue
        if spec is None and column in censored:
            # the distances measured, classed by the bands measured; the odd
            # value a little beyond the largest (a point along a long edge
            # from its node) is clipped into the top class, as any outlier is
            breaks[column] = {
                'kind': 'classes',
                'edges': [0.0] + [float(d) for d in censored[column]['bands']],
                'open_low': False,
                'open_high': False,
            }
            continue
        if spec is None:
            spec = MEASURE_BREAKS.get(by_measure.get(column))
        breaks[column] = class_breaks(
            described.get('units'),
            span,
            spec,
            targets.get(column),
        )
        if column in composite:
            breaks[column]['ramp'] = composite[column]['ramp']
            breaks[column]['centre'] = composite[column]['centre']
    return breaks


def censored_shares(stats, distances, access, weights):
    """Count the areas with nothing within the distance searched.

    An area whose distance is missing but whose access at the largest band is
    0 was measured, and nothing was found within reach: its population belongs
    in the chart, as a share of its own (``beyond``), rather than being
    dropped from it -- which would describe a development 1.5 km from any
    market by the few areas that are nearer.  The class shares are re-based to
    include it; areas with neither value remain missing.  Returns the updated
    summary (a new one where every area is beyond).
    """
    d = pd.to_numeric(distances, errors='coerce').to_numpy(dtype=float)
    a = pd.to_numeric(access, errors='coerce').to_numpy(dtype=float)
    w = np.asarray(weights, dtype=float)
    beyond = np.isnan(d) & (a == 0)
    valued = ~np.isnan(d)
    if not beyond.any():
        return stats
    if w[beyond | valued].sum() <= 0:
        w = np.ones(len(d))
    total = w[beyond | valued].sum()
    share = round(float(w[beyond].sum() / total * 100), 1)
    if not stats:
        return {
            'n': 0,
            'n_missing': int((~valued).sum()),
            'beyond': share,
            'n_beyond': int(beyond.sum()),
        }
    scale = float(w[valued].sum() / total) if total else 0.0
    for key in ('shares', 'class_shares'):
        if key in stats:
            stats[key] = [round(v * scale, 1) for v in stats[key]]
    stats['beyond'] = share
    stats['n_beyond'] = int(beyond.sum())
    return stats


def column_stats(values, weights, edges=None, breaks=None):
    """Weighted summary and histogram for one column, or None if all missing."""
    v = pd.to_numeric(values, errors='coerce').to_numpy(dtype=float)
    mask = ~np.isnan(v)
    n = int(mask.sum())
    if n == 0:
        return None
    w = weights[mask]
    if w.sum() <= 0:  # unpopulated areas (Condesa): fall back to equal weight
        w = np.ones(n)
    vv = v[mask]
    low, high = float(vv.min()), float(vv.max())
    if edges is None:
        edges = (
            np.linspace(low, high, HISTOGRAM_BINS + 1)
            if high > low
            else np.array([low, low + 1.0])
        )
    idx = np.clip(np.digitize(vv, edges[1:-1]), 0, len(edges) - 2)
    total = w.sum()
    shares = [
        round(float(w[idx == i].sum() / total * 100), 1)
        for i in range(len(edges) - 1)
    ]
    summary = {
        'n': n,
        'n_missing': int(len(v) - n),
        'min': _number(low),
        'max': _number(high),
        'mean': _number(float((w * vv).sum() / total)),
        'p25': _number(weighted_quantile(vv, w, 0.25)),
        'p50': _number(weighted_quantile(vv, w, 0.50)),
        'p75': _number(weighted_quantile(vv, w, 0.75)),
        'bins': [_number(e) for e in edges],
        'shares': shares,
    }
    if breaks:
        # the share of population in each *map class*, exactly.  This is what
        # the dashboard's histogram draws, so that its bars stand over the
        # legend's swatches: a class edge falling inside a uniform bin cannot be
        # recovered from `shares` afterwards.
        classes = class_index(vv, breaks)
        summary['class_shares'] = [
            round(float(w[classes == i].sum() / total * 100), 1)
            for i in range(class_count(breaks))
        ]
        unclassed = float(w[classes < 0].sum() / total * 100)
        if unclassed > 0.05:
            summary['unclassed'] = round(unclassed, 1)
    return summary


def band_distribution(frame, weights, columns):
    """Weighted share of each nested access band, plus a final 'none' class.

    Bands are nested -- access within 500 m implies access within 1000 m -- so
    an area is classed by the smallest band it reaches, and the remainder have
    access at none of them.
    """
    present = [c for c in columns if c in frame.columns]
    if not present:
        return None
    band = np.full(len(frame), len(present))
    for i, column in reversed(list(enumerate(present))):
        values = pd.to_numeric(frame[column], errors='coerce').fillna(-1)
        band[values.to_numpy(dtype=float) >= BAND_THRESHOLD] = i
    w = weights if weights.sum() > 0 else np.ones(len(frame))
    total = w.sum()
    return [
        round(float(w[band == i].sum() / total * 100), 1)
        for i in range(len(present) + 1)
    ]


def _banded_column_sets(indicators):
    """``{'<family>|<network>': [band columns, ascending]}`` for every banded measure."""
    sets = {}
    for family in indicators['families']:
        for key, measure in family['measures'].items():
            if measure.get('kind') != BANDED:
                continue
            for network, bands in (measure.get('networks') or {}).items():
                ordered = [
                    bands[d] for d in sorted(bands, key=lambda x: int(x))
                ]
                sets[f'{family["id"]}|{key}|{network}'] = ordered
    return sets


def scale_stats(
    r,
    entry,
    scale,
    indicators,
    band_sets,
    edges_by_column=None,
    breaks_by_column=None,
):
    """Weighted summaries for every column exported at one scale.

    Read through the same left join the layer was exported with, so the counts
    describe the areas actually on the map: an area that received no source
    units is one without a value, not one that does not exist.
    """
    resolved = canonical_columns(r, scale['table'])
    frame = r.get_df(
        scale_query(scale, resolved, entry['columns'], with_geom=False),
    )
    weights = (
        pd.to_numeric(frame['pop_est'], errors='coerce')
        .fillna(0)
        .to_numpy(
            dtype=float,
        )
        if 'pop_est' in frame.columns
        else np.ones(len(frame))
    )
    columns = {}
    censored = indicators.get('censored') or {}
    for column in entry['columns']:
        edges = (edges_by_column or {}).get(column)
        stats = column_stats(
            frame[column],
            weights,
            edges,
            (breaks_by_column or {}).get(column),
        )
        access = (censored.get(column) or {}).get('access')
        if access and access in frame.columns:
            stats = censored_shares(
                stats,
                frame[column],
                frame[access],
                weights,
            )
        if stats:
            columns[column] = stats
    bands = {}
    for key, ordered in band_sets.items():
        distribution = band_distribution(frame, weights, ordered)
        if distribution:
            bands[key] = distribution
    return {
        'weighted_by': 'pop_est' if weights.sum() > 0 else None,
        'areas': int(len(frame)),
        'columns': columns,
        'bands': bands,
    }


KDE_POINTS = 64


def weighted_kde(values, weights, grid):
    """A weighted Gaussian kernel density estimate, evaluated on ``grid``.

    The bandwidth is Silverman's rule of thumb, taken with the weighted
    standard deviation and interquartile range and the effective sample size
    of the weights (Kish), so that a few heavily weighted areas smooth as the
    few observations they are.  Returns densities, or None where fewer than
    two values are present.
    """
    v = pd.to_numeric(pd.Series(values), errors='coerce').to_numpy(dtype=float)
    w = np.asarray(weights, dtype=float)
    mask = ~np.isnan(v) & ~np.isnan(w) & (w > 0)
    v, w = v[mask], w[mask]
    if len(v) < 2:
        return None
    w = w / w.sum()
    mean = float((w * v).sum())
    sd = math.sqrt(float((w * (v - mean) ** 2).sum()))
    iqr = weighted_quantile(v, w, 0.75) - weighted_quantile(v, w, 0.25)
    spread = min(sd, iqr / 1.34) if iqr > 0 else sd
    grid = np.asarray(grid, dtype=float)
    if not spread > 0:
        # a single value: a narrow peak a twentieth of the axis wide
        spread = (grid[-1] - grid[0]) / 20 or 1.0
    n_effective = 1.0 / float((w**2).sum())
    bandwidth = 0.9 * spread * n_effective ** (-0.2)
    z = (grid[:, None] - v[None, :]) / bandwidth
    density = (np.exp(-0.5 * z**2) * w[None, :]).sum(axis=1)
    return density / (bandwidth * math.sqrt(2 * math.pi))


def region_distributions(r, config, regions, scale, columns):
    """Smoothed distributions of each column over each region's areas.

    The dashboard compares the regions' results for one scale -- the 100 m
    grid by default (``dashboard.distribution_scale``) -- as overlaid smoothed
    histograms.  An area belongs to a region where a point on its surface lies
    within the region's summary boundary.  Areas are weighted by population
    (``pop_est``) unless the region sets ``grid_weight``: ``{aggregation:
    <key>, per_unit: <people>}`` weights each area by the units of another
    aggregation it contains -- Condesa's lots, at an assumed 3.2 residents
    each, where the development is too new for the census to have counted.

    Returns ``{'x': {column: [low, high]}, 'regions': {region: {column:
    {density, n, weight, mean}}}}``, each column's densities evaluated at
    ``KDE_POINTS`` points spanning a range shared by every region.
    """
    tables = set(r.get_tables())
    resolved = canonical_columns(r, scale['table'])
    wanted = [c for c in columns if c in resolved]
    if not wanted:
        return None
    declared = config.get('regions') or {}
    frames = {}
    for key, region in regions.items():
        summary = region.get('summary_scale')
        boundary = f'agg_{_sql_key(summary)}' if summary else None
        if boundary not in tables:
            continue
        spec = (declared.get(key) or {}).get('grid_weight')
        if isinstance(spec, dict) and spec.get('aggregation'):
            units = f'agg_{_sql_key(spec["aggregation"])}'
            if units not in tables:
                print(f'  ! {key}: grid_weight aggregation {units} not found')
                continue
            per_unit = float(spec.get('per_unit', 1))
            weight = (
                f'(SELECT count(*) FROM {units} u WHERE '
                'ST_Intersects(ST_PointOnSurface(u.geom), q.geom)) '
                f'* {per_unit}'
            )
        elif 'pop_est' in resolved:
            weight = 'q.pop_est'
        else:
            weight = '1'
        context = ['pop_est'] if 'pop_est' in resolved else []
        inner = scale_query(
            scale,
            resolved,
            list(dict.fromkeys(wanted + context)),
        )
        select = ', '.join(f'q."{c}"' for c in wanted)
        frame = r.get_df(
            f'SELECT {select}, {weight} AS _weight FROM ({inner}) q '
            f'WHERE ST_Intersects(ST_PointOnSurface(q.geom), '
            f'(SELECT ST_Union(geom) FROM {boundary}))',
        )
        if frame is not None and len(frame):
            frames[key] = frame
    if not frames:
        return None
    x, out = {}, {key: {} for key in frames}
    for column in wanted:
        values = pd.concat(
            [
                pd.to_numeric(f[column], errors='coerce')
                for f in frames.values()
            ],
        ).dropna()
        if values.empty:
            continue
        low, high = float(values.min()), float(values.max())
        pad = (high - low) * 0.05 or 1.0
        low, high = low - pad, high + pad
        grid = np.linspace(low, high, KDE_POINTS)
        x[column] = [_number(low), _number(high)]
        for key, frame in frames.items():
            weights = (
                pd.to_numeric(frame['_weight'], errors='coerce')
                .fillna(0)
                .to_numpy(dtype=float)
            )
            if not weights.sum() > 0:
                weights = np.ones(len(frame))
            density = weighted_kde(frame[column], weights, grid)
            if density is None:
                continue
            present = pd.to_numeric(frame[column], errors='coerce').notna()
            w = weights[present.to_numpy()]
            v = pd.to_numeric(frame[column], errors='coerce')[present]
            out[key][column] = {
                'density': [float(f'{d:.4g}') for d in density],
                'n': int(present.sum()),
                'weight': _number(float(w.sum())),
                'mean': (
                    _number(float((v * w).sum() / w.sum()))
                    if w.sum() > 0
                    else None
                ),
            }
    return {'points': KDE_POINTS, 'x': x, 'regions': out}


# ---------------------------------------------------------------------------
# Manifest


def region_definitions(r, config, scale_entries):
    """The regions of interest a pane can be set to.

    Declared in the ``dashboard.regions`` block, because which custom
    aggregation summarises a region -- and which scales belong to it -- is a
    reporting decision, not something the aggregation config records.  With none
    declared the whole study region is the single region, at every scale.
    """
    available = {e['key'] for e in scale_entries}
    declared = config.get('regions') or {}
    if not declared:
        return {
            'study_region': {
                'label': _labels(config['label'], r.name),
                'summary_scale': None,
                'summary_table': 'indicators_region',
                'scales': [e['key'] for e in scale_entries],
            },
        }
    regions = {}
    for key, spec in declared.items():
        spec = spec or {}
        summary = spec.get('summary_scale')
        regions[key] = {
            'label': _labels(spec.get('label'), key),
            'summary_scale': summary,
            'summary_table': (
                f'indicators_{_sql_key(summary)}'
                if summary
                else 'indicators_region'
            ),
            'scales': [
                s
                for s in (spec.get('scales') or sorted(available))
                if s in available
            ],
        }
    return regions


def region_bounds(r, regions):
    """Each region's bounding box in EPSG:4326, for framing a pane.

    Condesa is a small part of the study area, so a pane showing it should not
    open on the whole of Mexicali.
    """
    tables = set(r.get_tables())
    for region in regions.values():
        scale = region.get('summary_scale')
        table = f'agg_{_sql_key(scale)}' if scale else 'urban_study_region'
        if table not in tables:
            continue
        gdf = r.get_gdf(f'SELECT geom FROM {table}')
        if gdf is None or len(gdf) == 0:
            continue
        region['bbox'] = [
            round(float(v), 5) for v in gdf.to_crs(4326).total_bounds
        ]
    return regions


def region_values(r, regions, columns):
    """The authoritative region-wide value of each indicator, per region.

    Only the region layer tiles the whole study area, so this -- not an average
    over the areas of whatever scale happens to be displayed -- is the figure
    the dashboard reports as the region's own.
    """
    values = {}
    tables = set(r.get_tables())
    for key, region in regions.items():
        table = region['summary_table']
        if table not in tables:
            continue
        resolved = canonical_columns(r, table)
        wanted = [c for c in columns if c in resolved]
        if not wanted:
            continue
        select = ', '.join(
            f'ROUND("{resolved[c][0]}"::numeric, 2)::float8 AS "{c}"'
            for c in wanted
        )
        frame = r.get_df(f'SELECT {select} FROM {table}')
        if len(frame) == 0:
            continue
        row = frame.iloc[0]
        values[key] = {c: _number(row[c]) for c in wanted}
    return values


def data_sources(r):
    """Attribution for the map image export and the info panel."""
    sources = []
    osm = r.config.get('OpenStreetMap') or {}
    if osm:
        sources.append(
            {
                'name': 'OpenStreetMap',
                'source': osm.get('source') or 'OpenStreetMap contributors',
                'publication_date': str(osm.get('publication_date') or ''),
                'url': osm.get('url') or 'https://www.openstreetmap.org',
                'licence': osm.get('licence') or 'ODbL',
            },
        )
    population = r.config.get('population') or {}
    if population:
        sources.append(
            {
                'name': population.get('name') or 'Population',
                'source': population.get('source') or '',
                'publication_date': str(
                    population.get('publication_date') or '',
                ),
                'url': population.get('url') or '',
                'licence': population.get('licence') or '',
                'citation': population.get('citation') or '',
            },
        )
    seen = set()
    for name, spec in (r.config.get('points_of_interest') or {}).items():
        citation = (spec or {}).get('citation') or (spec or {}).get('source')
        if not citation or citation in seen:
            continue
        seen.add(citation)
        sources.append(
            {
                'name': (spec or {}).get('dest_name_full') or name,
                'source': (spec or {}).get('source') or '',
                'publication_date': str(
                    (spec or {}).get('publication_date') or '',
                ),
                'url': (spec or {}).get('url') or '',
                'licence': (spec or {}).get('licence') or '',
                'citation': citation,
            },
        )
    return sources


def copy_data_dictionary(r, outdir):
    """Copy the region's generated data dictionary alongside the layers."""
    region_dir = r.config['region_dir']
    copied = {}
    for extension in ('csv', 'xlsx', 'pdf'):
        source = f'{region_dir}/{r.codename}_data_dictionary.{extension}'
        if os.path.exists(source):
            target = f'data_dictionary.{extension}'
            shutil.copyfile(source, f'{outdir}/{target}')
            copied[extension] = target
    if not copied:
        print(
            '  ! no data dictionary found; run generate() for this region '
            'to produce one',
        )
    return copied


def dictionary_descriptions(outdir, dictionary):
    """Descriptions from the generated data dictionary, keyed by variable.

    Preferred over ``describe_variable`` where present, because it is the text
    the region's own published dictionary uses, and it carries units and the
    statistic alongside.
    """
    if 'csv' not in dictionary:
        return {}
    out = {}
    with open(f'{outdir}/{dictionary["csv"]}', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            variable = (row.get('Variable') or '').strip()
            if not variable:
                continue
            out[variable] = {
                'category': (row.get('Category') or '').strip(),
                'en': (row.get('Indicator') or '').strip(),
                'units': (row.get('Units') or '').strip(),
                'statistic': (row.get('Statistic') or '').strip(),
            }
    return out


def attenuation_parameters(r):
    """What walkability's heat attenuation was computed with, for the viewer.

    ``{lambda, percentiles, heat: {letter: {variable, label, bounds}}}``,
    recorded with the walkability variants (see _walkability_variants), or
    None where none were computed or none attenuate.
    """
    try:
        from _walkability_variants import recorded_bounds, walkability_config

        config = walkability_config(r)
        if not config or 'multiplicative' not in config['forms']:
            return None
        recorded = recorded_bounds(r)
    except Exception as e:  # reported, and the dashboard exports without it
        print(f'  ! walkability attenuation parameters unavailable: {e}')
        return None
    if not recorded:
        return None
    heat = recorded.get('heat') or {}
    for letter, measure in heat.items():
        measure['label'] = _labels(measure.get('label'), letter)
    return {
        'lambda': recorded.get('attenuation'),
        'percentiles': recorded.get('percentiles'),
        'heat': heat,
    }


def export(
    r,
    outdir=None,
    only_scales=None,
    layers=True,
    kind=None,
):
    """Export dashboard layers, vocabulary and statistics for a region.

    Takes a loaded Region, or the codename or configuration path of one.
    ``kind`` is the type of dashboard, ``general`` or ``composite`` (see
    :func:`dashboard_type`), by default the configured ``dashboard.type``.
    """
    if type(r) is str:
        r = ghsci.Region(r)
    config = dashboard_config(r)
    kind, slug = dashboard_type(config, kind)
    outdir = outdir or f'/tmp/dashboard_export/{slug}'
    os.makedirs(outdir, exist_ok=True)
    print(f'{r.name} ({r.codename}) {kind} dashboard -> {outdir}', flush=True)

    scales = discover_scales(r, config)
    if only_scales:
        scales = [s for s in scales if s['key'] in only_scales]
    if not scales:
        sys.exit('No aggregation scales found for this region.')
    print(f'  Scales: {", ".join(s["key"] for s in scales)}')

    # the vocabulary is derived once, from the union of every scale's columns,
    # so a permutation missing at one scale is still offered where it exists
    available = set()
    for scale in scales:
        available.update(canonical_columns(r, scale['table']))
    indicators = build_indicators(r, config, available)
    select_dashboard(indicators, kind, config.get('index'))
    # where each distance's search stopped, so that a distance beyond it is
    # shown as such rather than as no data
    exported_columns = {
        c for f in indicators['families'] for c in f['columns']
    }
    indicators['censored'] = {
        k: v
        for k, v in censored_distances(r, available).items()
        if k in exported_columns
    }
    print(
        f'  Vocabulary: {len(indicators["families"])} families over '
        f'{len(indicators["descriptions"])} variables',
    )

    vocabulary_columns = sorted(
        {c for f in indicators['families'] for c in f['columns']},
    )
    # tiled in groups of bounded width, so that no feature is too heavy to
    # survive the zoomed out tiles; a composite dashboard is one group
    if kind == 'composite':
        name = indicators['families'][0]['id'].replace('composite_', '', 1)
        groups = {_sql_key(name) or 'index': vocabulary_columns}
    else:
        groups = tile_groups(
            indicators['families'],
            int(config.get('tile_max_columns') or TILE_MAX_COLUMNS),
        )
    indicators['column_group'] = {
        c: group for group, columns in groups.items() for c in columns
    }
    # the scales drawn as a smoothed surface on request: the regular grid
    smooth = set(config.get('smooth_scales') or ['grid'])
    # exported scales are kept paired with their definition: a scale that fails
    # to export is dropped from both, so the stats pass below cannot silently
    # pair one scale's columns with another scale's table
    exported = []
    for scale in scales:
        resolved_scale, entry = export_scale(
            r,
            scale,
            vocabulary_columns,
            outdir,
            write=layers,
            groups=groups,
            index_decimals=1 if kind in INDEX_TYPES else None,
            raster=scale['key'] in smooth,
        )
        if entry:
            exported.append((entry, resolved_scale))
    entries = [entry for entry, _ in exported]

    if layers:
        layer_index = export_context_layers(r, indicators, outdir)
    else:
        # --no-layers regenerates only the JSON, so keep the layer index that
        # the previous full export recorded rather than dropping it
        layer_index = _previous_layers(outdir)
    dictionary = copy_data_dictionary(r, outdir)
    published = dictionary_descriptions(outdir, dictionary)
    for variable, entry in published.items():
        current = indicators['descriptions'].get(variable)
        if not current:
            continue
        # the published text is the region's own wording, so it wins for prose;
        # units and statistic stay as describe_units() resolves them now
        for key, value in entry.items():
            if value and key not in ('units', 'statistic'):
                current[key] = value

    regions = region_bounds(r, region_definitions(r, config, entries))
    boundary = r.get_gdf(
        'SELECT geom FROM urban_study_region',
    ).to_crs(4326)
    manifest = {
        'slug': slug,
        'type': kind,
        'label': config['label'],
        'title': config['title'],
        'codename': r.codename,
        'name': r.name,
        'year': r.config.get('year'),
        'crs': r.config['crs']['srid'],
        'bbox': [round(float(v), 5) for v in boundary.total_bounds],
        'scales': {e['key']: e for e in entries},
        'layers': layer_index,
        # optional design tokens, applied by the viewer as CSS variables so a
        # region can be restyled without touching the site's code
        'style': config.get('style') or {},
        'regions': regions,
        'region_values': region_values(r, regions, vocabulary_columns),
        'data_dictionary': dictionary,
        'sources': data_sources(r),
        # the family the dashboard opens on: a composite dashboard's index,
        # whose profile is the featured view
        'featured': featured_family(config, indicators),
        # each language's conceptual model figure, where one is configured
        'conceptual_models': conceptual_models(r, outdir),
    }
    if kind in INDEX_TYPES:
        manifest['attenuation'] = attenuation_parameters(r)
    if not layers:
        # --no-layers keeps the rasters the last full export wrote
        previous = _previous_manifest(outdir)
        for key, entry in manifest['scales'].items():
            old = ((previous.get('scales') or {}).get(key) or {}).get('raster')
            if old:
                entry['raster'] = old

    band_sets = _banded_column_sets(indicators)
    ranges = column_ranges(r, exported)
    # one set of bin edges per column, shared by every scale
    edges_by_column = {
        column: histogram_edges(span) for column, span in ranges.items()
    }
    # ...and one set of colour classes, so that a class means the same thing on
    # both panes and the share of population in each can be counted exactly
    targets = resolve_targets(r, config)
    breaks_by_column = all_class_breaks(
        indicators,
        ranges,
        targets,
        config.get('breaks'),
        composite_spreads(r, exported, indicators),
    )
    indicators['breaks'] = breaks_by_column
    indicators['targets'] = {
        column: target
        for column, target in targets.items()
        if column in breaks_by_column
    }
    print(
        f'  Classification: {len(breaks_by_column)} columns, '
        f'{len(indicators["targets"])} with a target',
    )
    stats = {}
    for entry, scale in exported:
        print(f'  distributions: {entry["key"]}', flush=True)
        stats[entry['key']] = scale_stats(
            r,
            entry,
            scale,
            indicators,
            band_sets,
            edges_by_column,
            breaks_by_column,
        )

    outputs = [
        ('manifest.json', manifest),
        ('indicators.json', indicators),
        ('stats.json', stats),
    ]
    if kind in INDEX_TYPES:
        # each region's results, as smoothed distributions for comparing them
        key = config.get('distribution_scale') or 'grid'
        scale = next((s for e, s in exported if e['key'] == key), None)
        scores = [c for c in vocabulary_columns if c.startswith('index_')]
        # and each indicator's own values, in its natural units (metres, per
        # cent, degrees), which the settings show in place of its score
        for family in indicators['families']:
            structure = family.get('composite') or {}
            for indicator in structure.get('indicators') or []:
                if indicator.get('area_column'):
                    scores.append(indicator['area_column'])
            for variant in structure.get('variants') or []:
                scores += list((variant.get('area_columns') or {}).values())
        scores = list(dict.fromkeys(scores))
        distributions = (
            region_distributions(r, config, regions, scale, scores)
            if scale
            else None
        )
        if distributions:
            distributions['scale'] = key
            manifest['distributions'] = 'distributions.json'
            outputs.append(('distributions.json', distributions))
            print(
                f'  distributions: {len(distributions["x"])} columns for '
                f'{", ".join(distributions["regions"])} by {key}',
            )
    for name, payload in outputs:
        with open(f'{outdir}/{name}', 'w', encoding='utf-8') as f:
            json.dump(payload, f, ensure_ascii=False, separators=(',', ':'))
        size = os.path.getsize(f'{outdir}/{name}') / 1024
        print(f'  {name}: {size:,.0f} KB')
    print(f'  bbox {manifest["bbox"]}', flush=True)
    return manifest


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    flags = [a for a in sys.argv[1:] if a.startswith('--')]
    if not args:
        sys.exit(__doc__)
    only = None
    kind = None
    for flag in flags:
        if flag.startswith('--scales='):
            only = [s.strip() for s in flag.split('=', 1)[1].split(',')]
        if flag.startswith('--type='):
            kind = flag.split('=', 1)[1].strip()
    # --no-layers regenerates the manifest, vocabulary and statistics without
    # rewriting a gigabyte of GeoJSONSeq: revising labels, themes or
    # interventions should not cost a full re-export and re-tile
    export(
        args[0],
        args[1] if len(args) > 1 else None,
        only,
        layers='--no-layers' not in flags,
        kind=kind,
    )
